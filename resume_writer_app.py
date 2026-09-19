from __future__ import annotations

import os
import re
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Iterable, Optional, Tuple

os.environ.setdefault("TK_SILENCE_DEPRECATION", "1")

import tkinter as tk
from tkinter import filedialog, messagebox

from formats_store import DEFAULT_FORMAT_VALUES, get_format_for_writer
from user_auth import get_current_user


AUTHOR_NAME = "Yashashchandra Kollu"
EMAIL = "yashashchandrakollu1@gmail.com"
PHONE = "+1 (205) 897 7790"
PHONE_LINK = "tel:+1%20(205)%20897%207790"
LOCATION = "Atlanta, GA"
APP_VERSION = "v2.1.0"
DEFAULT_OUTPUT_DIR = Path("/Users/yck/Desktop/CLGENAPPL")

FONT_CHOICES = (
    "Aptos",
    "Arial",
    "Calibri",
    "Cambria",
    "Garamond",
    "Georgia",
    "Times New Roman",
    "Verdana",
)
NAME_SIZE_CHOICES = ("14", "15", "16", "17", "18", "20")
HEADING_SIZE_CHOICES = ("11", "12", "13", "14")
BODY_SIZE_CHOICES = ("10", "11", "12")
MONTH_PATTERN = (
    r"Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?|"
    r"Jul(?:y)?|Aug(?:ust)?|Sep(?:tember)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?"
)
DATE_RANGE_RE = re.compile(
    rf"\b(?:{MONTH_PATTERN})\s+\d{{4}}\s*(?:-|–|—|to)\s*(?:Present|Current|(?:{MONTH_PATTERN})\s+\d{{4}})\b",
    re.IGNORECASE,
)
COMPANY_DATE_DEFAULTS = {
    "first_horizon": "August 2025 – Present",
    "dbs": "July 2022 – July 2024",
    "value_labs": "February 2021 – March 2022",
}
COMPANY_ALIASES = {
    "first_horizon": ("first horizon bank",),
    "dbs": ("dbs bank", "dbs tech india"),
    "value_labs": ("value labs",),
}
ROLE_DATE_DEFAULTS_BY_COMPANY = {
    "first_horizon": ("August 2025 – Present",),
    "dbs": ("September 2023 – July 2024", "July 2022 – September 2023"),
    "value_labs": ("February 2021 – March 2022",),
}

THEMES = {
    "light": {
        "app_bg": "#f5f5f7",
        "panel_bg": "#ffffff",
        "text_bg": "#ffffff",
        "text_fg": "#1d1d1f",
        "muted_fg": "#6e6e73",
        "accent": "#007aff",
        "border": "#d2d2d7",
        "button_bg": "#f5f5f7",
        "button_active": "#e8e8ed",
        "disabled_bg": "#8e8e93",
        "disabled_fg": "#ffffff",
        "error_fg": "#b91c1c",
        "ok_fg": "#166534",
        "select_bg": "#acd7ff",
        "toggle_track": "#d1d1d6",
        "toggle_knob": "#ffffff",
        "toggle_text": "#1d1d1f",
    },
    "dark": {
        "app_bg": "#000000",
        "panel_bg": "#1c1c1e",
        "text_bg": "#1c1c1e",
        "text_fg": "#f5f5f7",
        "muted_fg": "#a1a1a6",
        "accent": "#0a84ff",
        "border": "#38383a",
        "button_bg": "#2c2c2e",
        "button_active": "#3a3a3c",
        "disabled_bg": "#2c2c2e",
        "disabled_fg": "#636366",
        "error_fg": "#fca5a5",
        "ok_fg": "#86efac",
        "select_bg": "#005ecb",
        "toggle_track": "#30d158",
        "toggle_knob": "#ffffff",
        "toggle_text": "#f5f5f7",
    },
}


DEFAULT_MASTERS_EDUCATION = (
    "Master of Science in Computer Science\tAugust 2024 – May 2026\n"
    "University of Alabama at Birmingham, AL  |  GPA: 4.0 / 4.0"
)
DEFAULT_BACHELORS_EDUCATION = (
    "Bachelor of Technology in Information Technology\tAugust 2018 – July 2022\n"
    "Institute of Aeronautical Engineering, Hyderabad, India | GPA: 3.5 / 4.0"
)


def default_output_filename() -> str:
    today = date.today()
    return f"YcKResumeXXX{today:%d%b}.docx"


def default_output_filename_for(company_part: str = "") -> str:
    today = date.today()
    clean_part = re.sub(r"[^A-Za-z0-9_-]+", "", company_part.strip())
    return f"YcKResume{clean_part or 'XXX'}{today:%d%b}.docx"


def find_matching_output_files(output_dir: Path, name_part: str, *, include_applied: bool = True) -> list[Path]:
    clean_part = re.sub(r"[^A-Za-z0-9_-]+", "", name_part.strip())
    if not clean_part or not output_dir.exists() or not output_dir.is_dir():
        return []
    needle = clean_part.casefold()
    matches = []
    search_dirs = [output_dir]
    applied_dir = output_dir / "#applied"
    if include_applied and applied_dir.exists() and applied_dir.is_dir():
        search_dirs.append(applied_dir)
    for search_dir in search_dirs:
        for path in search_dir.iterdir():
            if path.suffix.casefold() not in {".docx", ".pdf"}:
                continue
            if needle in path.name.casefold():
                matches.append(path)
    return sorted(matches, key=lambda path: path.name.casefold())


def load_docx_dependencies() -> None:
    global Document, Inches, OxmlElement, Pt, WD_ALIGN_PARAGRAPH, qn

    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt
    except ImportError as exc:
        raise RuntimeError("python-docx is required. Run: python3 -m pip install -r requirements.txt") from exc


@dataclass(frozen=True)
class ResumeFormat:
    font_name: str = DEFAULT_FORMAT_VALUES["font_name"]
    name_size: int = DEFAULT_FORMAT_VALUES["name_size"]
    heading_size: int = DEFAULT_FORMAT_VALUES["heading_size"]
    body_size: int = DEFAULT_FORMAT_VALUES["body_size"]


def resume_format_from_store(username: Optional[str] = None) -> ResumeFormat:
    """Build a ResumeFormat from the user's primary format (or app default).

    Safe fallback: any store/auth error yields the protected default values.
    The Formatter UI still controls generation via its own widgets; App can call
    this when wiring login / primary-format selection into generate.
    """
    try:
        spec = get_format_for_writer(username if username is not None else get_current_user())
        return ResumeFormat(**spec.to_writer_kwargs())
    except Exception:
        return ResumeFormat()


@dataclass(frozen=True)
class ResumeContent:
    name: str
    email: str
    phone: str
    location: str
    summary: str
    skills: str
    experience: str
    certifications: str
    top_skills: str
    education_entries: Tuple[str, ...]


def normalize_experience_bullet_text(text: str) -> str:
    return re.sub(r"\s+—\s+", ", ", text.strip())


def clean_lines(text: str) -> list[str]:
    return [line.rstrip() for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n")]


def split_keywords(text: str) -> list[str]:
    parts = re.split(r"[,;\n]", text)
    return [part.strip() for part in parts if part.strip()][:5]


def count_words(text: str) -> int:
    return len(re.findall(r"\b[\w'-]+\b", text))


def phone_hyperlink(phone: str) -> str:
    compact = re.sub(r"\s+", "%20", phone.strip())
    return f"tel:{compact}"


def normalize_date_text(text: str) -> str:
    return re.sub(r"\s*(?:-|–|—|to)\s*", " – ", text.strip(), flags=re.IGNORECASE)


def normalize_dated_left_text(text: str) -> str:
    normalized = re.sub(r"\s+", " ", text.strip())
    return re.sub(r"\s*\|\s*", " | ", normalized).strip(" |,-")


def split_right_side_date(left: str, right: str) -> Optional[Tuple[str, str]]:
    match = DATE_RANGE_RE.search(right)
    if not match:
        return None
    right_prefix = normalize_dated_left_text(right[: match.start()])
    left_parts = [normalize_dated_left_text(left)]
    if right_prefix:
        left_parts.append(right_prefix)
    return " ".join(part for part in left_parts if part), normalize_date_text(match.group(0))


def split_dated_line(line: str) -> Optional[Tuple[str, str]]:
    stripped = line.strip()

    if "\t" in stripped:
        left, right = stripped.rsplit("\t", 1)
        dated = split_right_side_date(left, right)
        if dated:
            return dated

    if "|" in stripped:
        left, right = stripped.rsplit("|", 1)
        dated = split_right_side_date(left, right)
        if dated:
            return dated

    match = DATE_RANGE_RE.search(stripped)
    if match and match.start() > 0:
        left = normalize_dated_left_text(stripped[: match.start()])
        if left:
            return left, normalize_date_text(match.group(0))

    return None


def detect_known_company(line: str) -> Optional[str]:
    normalized = re.sub(r"\s+", " ", line.strip())
    normalized_lower = normalized.lower()
    if DATE_RANGE_RE.search(normalized):
        return None
    for company_key, aliases in COMPANY_ALIASES.items():
        if any(normalized_lower.startswith(alias) for alias in aliases):
            return company_key
    return None


def split_known_company_line(line: str) -> Optional[Tuple[str, str, str]]:
    company_key = detect_known_company(line)
    if not company_key:
        return None
    left = re.sub(r"\s+", " ", line.strip()).strip(" |,-")
    left = re.sub(r"\s*\|\s*", " | ", left)
    return left, COMPANY_DATE_DEFAULTS[company_key], company_key


def fallback_role_date(company_key: Optional[str], role_index: int) -> Optional[str]:
    if not company_key:
        return None
    dates = ROLE_DATE_DEFAULTS_BY_COMPANY.get(company_key, ())
    if not dates:
        return None
    return dates[min(role_index, len(dates) - 1)]


def looks_like_role_line(line: str) -> bool:
    stripped = line.strip()
    if not stripped or stripped.endswith("."):
        return False
    if len(stripped.split()) > 10:
        return False
    if re.search(r"\b(built|created|designed|developed|implemented|improved|reduced|delivered|managed|maintained)\b", stripped, re.IGNORECASE):
        return False
    if re.search(r"\b(engineer|developer|analyst|architect|consultant|intern|lead|manager|specialist|associate)\b", stripped, re.IGNORECASE):
        return True
    return False


def set_run_font(run, font_name: str, size: int, bold: bool = False) -> None:
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    run.font.size = Pt(size)
    run.bold = bold


def format_paragraph(paragraph, before: float = 0, after: float = 3, alignment=None) -> None:
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.0
    if alignment is not None:
        paragraph.alignment = alignment


def add_border_rule(document: Document) -> None:
    paragraph = document.add_paragraph()
    format_paragraph(paragraph, before=0, after=3)
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "2")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "AAAAAA")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def add_hyperlink(paragraph, text: str, url: str, fmt: ResumeFormat, bold: bool = False) -> None:
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")

    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), fmt.font_name)
    fonts.set(qn("w:hAnsi"), fmt.font_name)
    fonts.set(qn("w:eastAsia"), fmt.font_name)
    r_pr.append(fonts)

    size = OxmlElement("w:sz")
    size.set(qn("w:val"), str(fmt.body_size * 2))
    r_pr.append(size)

    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    r_pr.append(color)

    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)

    if bold:
        r_pr.append(OxmlElement("w:b"))

    new_run.append(r_pr)
    text_element = OxmlElement("w:t")
    text_element.text = text
    new_run.append(text_element)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_text_run(paragraph, text: str, fmt: ResumeFormat, bold: bool = False) -> None:
    run = paragraph.add_run(text)
    set_run_font(run, fmt.font_name, fmt.body_size, bold=bold)


def add_section_heading(document: Document, title: str, fmt: ResumeFormat) -> None:
    paragraph = document.add_paragraph()
    format_paragraph(paragraph, before=8, after=3, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    run = paragraph.add_run(title)
    set_run_font(run, fmt.font_name, fmt.heading_size, bold=True)


def add_plain_paragraph(
    document: Document,
    text: str,
    fmt: ResumeFormat,
    alignment=None,
    before: float = 2,
    after: float = 4,
) -> None:
    paragraph = document.add_paragraph()
    if alignment is None:
        alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    format_paragraph(paragraph, before=before, after=after, alignment=alignment)
    add_text_run(paragraph, text, fmt)


def add_skill_line(document: Document, line: str, fmt: ResumeFormat) -> None:
    paragraph = document.add_paragraph()
    format_paragraph(paragraph, before=2.5, after=2.5, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    if ":" in line:
        label, rest = line.split(":", 1)
        add_text_run(paragraph, f"{label.strip()}: ", fmt, bold=True)
        add_text_run(paragraph, rest.strip(), fmt)
    else:
        add_text_run(paragraph, line, fmt)


def add_bullet(document: Document, text: str, fmt: ResumeFormat, *, normalize_dash: bool = False) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    format_paragraph(paragraph, before=2, after=2, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    if normalize_dash:
        text = normalize_experience_bullet_text(text)
    add_text_run(paragraph, text, fmt)


def set_cell_width(cell, width_inches: float) -> None:
    cell.width = Inches(width_inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(round(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_no_wrap(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    no_wrap = tc_pr.first_child_found_in("w:noWrap")
    if no_wrap is None:
        tc_pr.append(OxmlElement("w:noWrap"))


def set_table_grid(table, widths_inches: tuple[float, ...]) -> None:
    tbl_grid = table._tbl.tblGrid
    for grid_col in list(tbl_grid):
        tbl_grid.remove(grid_col)
    for width_inches in widths_inches:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(round(width_inches * 1440)))
        tbl_grid.append(grid_col)


def set_cell_margins(cell, *, top: int = 0, start: int = 0, bottom: int = 0, end: int = 0) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        margin = tc_mar.find(qn(f"w:{margin_name}"))
        if margin is None:
            margin = OxmlElement(f"w:{margin_name}")
            tc_mar.append(margin)
        margin.set(qn("w:w"), str(value))
        margin.set(qn("w:type"), "dxa")


def remove_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for border_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{border_name}"))
        if border is None:
            border = OxmlElement(f"w:{border_name}")
            borders.append(border)
        border.set(qn("w:val"), "nil")


def show_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for border_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{border_name}"))
        if border is None:
            border = OxmlElement(f"w:{border_name}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "6")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "FF0000")


def add_tabbed_line(document: Document, left_text: str, right_text: str, fmt: ResumeFormat) -> None:
    right_text = normalize_date_text(right_text) if DATE_RANGE_RE.search(right_text) else right_text.strip()
    table = document.add_table(rows=1, cols=2)
    table.autofit = False
    set_table_grid(table, (4.45, 2.55))
    remove_table_borders(table)

    left_cell, right_cell = table.rows[0].cells
    set_cell_width(left_cell, 4.45)
    set_cell_width(right_cell, 2.55)
    set_cell_margins(left_cell)
    set_cell_margins(right_cell)
    set_cell_no_wrap(right_cell)

    left_paragraph = left_cell.paragraphs[0]
    format_paragraph(left_paragraph, before=4, after=2, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    add_text_run(left_paragraph, left_text.strip(), fmt, bold=True)

    right_paragraph = right_cell.paragraphs[0]
    format_paragraph(right_paragraph, before=4, after=2, alignment=WD_ALIGN_PARAGRAPH.RIGHT)
    if right_text.strip():
        add_text_run(right_paragraph, right_text.strip(), fmt, bold=True)


def strip_bullet_marker(line: str) -> str:
    return re.sub(r"^(?:[-*]|\u2022)\s+", "", line.strip())


def add_experience_section(document: Document, lines: Iterable[str], fmt: ResumeFormat) -> None:
    current_company = None
    role_index_by_company: dict[str, int] = {}

    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue

        clean_line = strip_bullet_marker(stripped)
        dated = split_dated_line(clean_line)
        if dated:
            left, right = dated
            add_tabbed_line(document, left, right, fmt)
            detected_company = detect_known_company(left)
            if detected_company:
                current_company = detected_company
                role_index_by_company.setdefault(current_company, 0)
            elif looks_like_role_line(left):
                role_index = role_index_by_company.get(current_company or "", 0)
                if fallback_role_date(current_company, role_index):
                    role_index_by_company[current_company] = role_index + 1
            continue

        known_company = split_known_company_line(clean_line)
        if known_company:
            left, right, company_key = known_company
            add_tabbed_line(document, left, right, fmt)
            current_company = company_key
            role_index_by_company.setdefault(current_company, 0)
            continue

        if looks_like_role_line(clean_line):
            role_index = role_index_by_company.get(current_company or "", 0)
            role_date = fallback_role_date(current_company, role_index)
            if role_date:
                add_tabbed_line(document, clean_line, role_date, fmt)
                role_index_by_company[current_company] = role_index + 1
                continue

        add_bullet(document, clean_line, fmt, normalize_dash=True)


def add_multiline_section(
    document: Document,
    lines: Iterable[str],
    fmt: ResumeFormat,
    *,
    skill_mode: bool = False,
    paragraph_alignment=None,
) -> None:
    if paragraph_alignment is None:
        paragraph_alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
        bullet_match = re.match(r"^(?:[-*]|\u2022)\s+(.*)$", stripped)
        if bullet_match:
            add_bullet(document, bullet_match.group(1).strip(), fmt)
        elif skill_mode:
            add_skill_line(document, stripped, fmt)
        elif "\t" in stripped:
            left, right = stripped.split("\t", 1)
            add_tabbed_line(document, left, right, fmt)
        elif split_dated_line(stripped):
            left, right = split_dated_line(stripped)
            add_tabbed_line(document, left, right, fmt)
        else:
            add_plain_paragraph(document, stripped, fmt, alignment=paragraph_alignment, before=2, after=3)


def add_education_entry(document: Document, education_text: str, fmt: ResumeFormat) -> None:
    for line in clean_lines(education_text):
        stripped = line.strip()
        if not stripped:
            continue
        dated = split_dated_line(stripped)
        if dated:
            left, right = dated
            add_tabbed_line(document, left, right, fmt)
        else:
            add_plain_paragraph(document, stripped, fmt, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=3)


def apply_document_defaults(document: Document, fmt: ResumeFormat) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.header_distance = Inches(0.4917)
    section.footer_distance = Inches(0.4917)

    normal = document.styles["Normal"]
    normal.font.name = fmt.font_name
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), fmt.font_name)
    normal.font.size = Pt(fmt.body_size)
    normal.paragraph_format.line_spacing = 1.0

    bullet = document.styles["List Bullet"]
    bullet.font.name = fmt.font_name
    bullet._element.rPr.rFonts.set(qn("w:eastAsia"), fmt.font_name)
    bullet.font.size = Pt(fmt.body_size)


def build_resume(content: ResumeContent, fmt: ResumeFormat, output_path: Path) -> Path:
    load_docx_dependencies()
    document = Document()
    apply_document_defaults(document, fmt)

    core = document.core_properties
    core.author = content.name or AUTHOR_NAME
    core.subject = "my resume"
    core.keywords = ", ".join(split_keywords(content.top_skills))
    core.title = f"{content.name or AUTHOR_NAME} Resume"

    name_paragraph = document.add_paragraph()
    format_paragraph(name_paragraph, before=0, after=3, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    name_run = name_paragraph.add_run(content.name or AUTHOR_NAME)
    set_run_font(name_run, fmt.font_name, fmt.name_size, bold=True)

    contact = document.add_paragraph()
    format_paragraph(contact, before=0, after=3, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    email = content.email or EMAIL
    phone = content.phone or PHONE
    location = content.location or LOCATION
    add_hyperlink(contact, email, f"mailto:{email}", fmt)
    add_text_run(contact, "  |  ", fmt, bold=True)
    add_hyperlink(contact, phone, phone_hyperlink(phone), fmt)
    add_text_run(contact, "  |  ", fmt, bold=True)
    add_text_run(contact, location, fmt)

    add_border_rule(document)
    add_section_heading(document, "SUMMARY", fmt)
    add_multiline_section(document, clean_lines(content.summary), fmt, paragraph_alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)

    add_border_rule(document)
    add_section_heading(document, "SKILLS", fmt)
    add_multiline_section(
        document,
        clean_lines(content.skills),
        fmt,
        skill_mode=True,
        paragraph_alignment=WD_ALIGN_PARAGRAPH.LEFT,
    )

    add_border_rule(document)
    add_section_heading(document, "WORK EXPERIENCE", fmt)
    add_experience_section(document, clean_lines(content.experience), fmt)

    if content.education_entries:
        add_border_rule(document)
        add_section_heading(document, "EDUCATION", fmt)
        for education_entry in content.education_entries:
            add_education_entry(document, education_entry, fmt)

    if content.certifications.strip():
        add_border_rule(document)
        add_section_heading(document, "Certifications", fmt)
        add_multiline_section(
            document,
            clean_lines(content.certifications),
            fmt,
            paragraph_alignment=WD_ALIGN_PARAGRAPH.LEFT,
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path


class ResumeWriterApp(tk.Tk):
    def __init__(self) -> None:
        super().__init__()
        self.title(f"Resume Writer {APP_VERSION}")
        self.geometry("1060x880")
        self.minsize(860, 700)

        self.dark_mode_var = tk.BooleanVar(value=True)
        self.font_var = tk.StringVar(value=str(DEFAULT_FORMAT_VALUES["font_name"]))
        self.name_size_var = tk.StringVar(value=str(DEFAULT_FORMAT_VALUES["name_size"]))
        self.heading_size_var = tk.StringVar(value=str(DEFAULT_FORMAT_VALUES["heading_size"]))
        self.body_size_var = tk.StringVar(value=str(DEFAULT_FORMAT_VALUES["body_size"]))
        self.output_dir_var = tk.StringVar(value=str(DEFAULT_OUTPUT_DIR))
        self.output_name_part_var = tk.StringVar(value="")
        self.output_var = tk.StringVar(value=str(DEFAULT_OUTPUT_DIR / default_output_filename_for()))
        self.output_warning_var = tk.StringVar(value="")
        self.name_var = tk.StringVar(value=AUTHOR_NAME)
        self.email_var = tk.StringVar(value=EMAIL)
        self.phone_var = tk.StringVar(value=PHONE)
        self.location_var = tk.StringVar(value=LOCATION)
        self.masters_var = tk.BooleanVar(value=True)
        self.bachelors_var = tk.BooleanVar(value=False)
        self.summary_count_var = tk.StringVar(value="Summary: 0 words (required: 41-50)")
        self.validation_var = tk.StringVar(value="")

        self._widgets_by_role = {}
        self._submit_widgets = []
        self._suppress_output_path_refresh = False

        self._configure_style()
        self._build_menu()
        self._build_ui()
        self._bind_shortcuts()
        self._wire_validation()
        self._update_validation_state()

    def _configure_style(self) -> None:
        self.theme = THEMES["dark" if self.dark_mode_var.get() else "light"]
        self.option_add("*Font", "Arial 12")
        self.option_add("*Background", self.c("app_bg"))
        self.option_add("*Foreground", self.c("text_fg"))
        self.option_add("*Entry.Background", self.c("text_bg"))
        self.option_add("*Entry.Foreground", self.c("text_fg"))
        self.option_add("*Text.Background", self.c("text_bg"))
        self.option_add("*Text.Foreground", self.c("text_fg"))
        self.option_add("*insertBackground", self.c("text_fg"))
        self.configure(bg=self.c("app_bg"))

    def c(self, key: str) -> str:
        return self.theme[key]

    def _build_menu(self) -> None:
        menu_bar = tk.Menu(self)
        help_menu = tk.Menu(menu_bar, tearoff=False)
        help_menu.add_command(label="Usage Help", command=self._show_usage_help)
        help_menu.add_separator()
        help_menu.add_command(label="About Resume Writer", command=self._show_about)
        menu_bar.add_cascade(label="Help", menu=help_menu)
        self.config(menu=menu_bar)

    def _show_about(self) -> None:
        messagebox.showinfo(
            "About Resume Writer",
            f"Resume Writer\n{APP_VERSION}",
        )

    def _show_usage_help(self) -> None:
        messagebox.showinfo(
            "Usage Help",
            (
                "Required fields: Summary, Skills, Job Experience, Top 5 Skills.\n"
                "Summary must be 41-50 words.\n\n"
                "Tab moves to the next field.\n"
                "Generate with Command+Return on Mac or Ctrl+Enter on Windows/Linux.\n\n"
                f"Default output folder:\n{DEFAULT_OUTPUT_DIR}"
            ),
        )

    def _build_ui(self) -> None:
        root = self._track(tk.Frame(self, bg=self.c("app_bg"), padx=16, pady=16), "app_frame")
        root.pack(fill=tk.BOTH, expand=True)
        root.columnconfigure(0, weight=1)
        root.rowconfigure(1, weight=1)

        controls = self._section(root, "Format Options")
        controls.grid(row=0, column=0, sticky="ew")
        for column in range(10):
            controls.columnconfigure(column, weight=1 if column in (1, 3, 5, 7) else 0)

        self._combo(controls, "Font", self.font_var, FONT_CHOICES, 0, 0)
        self._combo(controls, "Name Size", self.name_size_var, NAME_SIZE_CHOICES, 0, 2)
        self._combo(controls, "Heading Size", self.heading_size_var, HEADING_SIZE_CHOICES, 0, 4)
        self._combo(controls, "Text Size", self.body_size_var, BODY_SIZE_CHOICES, 0, 6)
        self.theme_toggle = self._toggle_button(controls, self.dark_mode_var, self._toggle_theme)
        self.theme_toggle.grid(row=0, column=8, columnspan=2, sticky="e")

        text_area = self._track(tk.Frame(root, bg=self.c("app_bg")), "app_frame")
        text_area.grid(row=1, column=0, sticky="nsew", pady=(12, 12))
        text_area.columnconfigure(0, weight=1)
        text_area.columnconfigure(1, weight=1)
        text_area.rowconfigure(1, weight=1, minsize=96)
        text_area.rowconfigure(4, weight=1, minsize=96)
        text_area.rowconfigure(6, weight=1, minsize=96)

        self.summary_text = self._text_box(text_area, "Summary", 0, 0, height=7)
        self.skills_text = self._text_box(text_area, "Skills", 0, 1, height=7)
        summary_footer = self._label(text_area, "", muted=True)
        summary_footer.configure(textvariable=self.summary_count_var)
        summary_footer.grid(row=2, column=0, sticky="w", pady=(0, 8))
        self.experience_text = self._text_box(text_area, "Job Experience", 3, 0, columnspan=2, height=7)
        self.certifications_text = self._text_box(text_area, "Certifications", 5, 0, height=7)
        self.top_skills_text = self._text_box(text_area, "Top 5 Skills For Metadata", 5, 1, height=7)

        output = self._section(root, "Output")
        output.grid(row=2, column=0, sticky="ew")
        output.columnconfigure(1, weight=1)

        self._label(output, "File Path").grid(row=0, column=0, sticky="w", padx=(0, 8), pady=(0, 8))
        output_path_frame = self._track(tk.Frame(output, bg=self.c("panel_bg")), "panel_frame")
        output_path_frame.grid(row=0, column=1, columnspan=3, sticky="w", pady=(0, 8))
        output_path_label = self._label(output_path_frame, "", muted=True)
        output_path_label.configure(textvariable=self.output_var)
        output_path_label.pack(side=tk.LEFT)
        self._button(output_path_frame, "Browse", self._choose_output).pack(side=tk.LEFT)
        self._label(output, "File Name").grid(row=1, column=0, sticky="w", padx=(0, 8))
        file_name_frame = self._track(tk.Frame(output, bg=self.c("panel_bg")), "panel_frame")
        file_name_frame.grid(row=1, column=1, sticky="w")
        self.file_name_entry = self._entry(file_name_frame, self.output_name_part_var, width=25)
        self.file_name_entry.pack(side=tk.LEFT)
        self.file_name_check_label = self._label(file_name_frame, "")
        self.file_name_check_label.configure(font=("Arial", 14, "bold"), fg=self.c("ok_fg"))
        self.file_name_check_label.pack(side=tk.LEFT)
        self.generate_button = self._button(output, "Generate DOCX", self._generate, accent=True)
        self.generate_button.grid(row=1, column=3)
        self._submit_widgets.append(self.generate_button)
        self.output_warning_label = self._label(output, "", muted=True)
        self.output_warning_label.configure(textvariable=self.output_warning_var)
        self.output_warning_label.grid(row=2, column=1, columnspan=3, sticky="w", pady=(6, 0))

        header = self._section(root, "Header And Education")
        header.grid(row=3, column=0, sticky="ew", pady=(12, 0))
        for column in range(4):
            header.columnconfigure(column, weight=1)

        self._entry(header, self.name_var).grid(row=0, column=0, sticky="ew", padx=(0, 8), pady=(0, 8))
        self._entry(header, self.email_var).grid(row=0, column=1, sticky="ew", padx=(0, 8), pady=(0, 8))
        self._entry(header, self.phone_var).grid(row=0, column=2, sticky="ew", padx=(0, 8), pady=(0, 8))
        self._entry(header, self.location_var).grid(row=0, column=3, sticky="ew", pady=(0, 8))

        education_controls = self._track(tk.Frame(header, bg=self.c("panel_bg")), "panel_frame")
        education_controls.grid(row=1, column=0, columnspan=4, sticky="w")
        self.masters_choice = self._choice_label(education_controls, "Masters", self.masters_var, self._toggle_masters)
        self.masters_choice.pack(side=tk.LEFT)
        self._icon_button(education_controls, "✎", lambda: self._show_education_editor("masters")).pack(side=tk.LEFT, padx=(4, 18))
        self.bachelors_choice = self._choice_label(education_controls, "Bachelors", self.bachelors_var, self._toggle_bachelors)
        self.bachelors_choice.pack(side=tk.LEFT)
        self._icon_button(education_controls, "✎", lambda: self._show_education_editor("bachelors")).pack(side=tk.LEFT, padx=(4, 0))

        self.education_editor_frame = self._track(tk.Frame(header, bg=self.c("panel_bg")), "panel_frame")
        self.education_editor_frame.grid(row=2, column=0, columnspan=4, sticky="ew", pady=(8, 0))
        self.education_editor_frame.columnconfigure(0, weight=1)
        self.education_editor_label = self._label(self.education_editor_frame, "Masters Education")
        self.education_editor_label.grid(row=0, column=0, sticky="w")
        self.masters_education_text = self._text_widget(self.education_editor_frame, height=3)
        self.bachelors_education_text = self._text_widget(self.education_editor_frame, height=3)
        self.masters_education_text.insert("1.0", DEFAULT_MASTERS_EDUCATION)
        self.bachelors_education_text.insert("1.0", DEFAULT_BACHELORS_EDUCATION)
        self.masters_education_text.grid(row=1, column=0, sticky="ew", pady=(4, 0))
        self.bachelors_education_text.grid(row=1, column=0, sticky="ew", pady=(4, 0))
        self.education_editor_frame.grid_remove()
        self.bachelors_education_text.grid_remove()

    def _section(self, parent, label: str) -> tk.LabelFrame:
        return self._track(tk.LabelFrame(
            parent,
            text=label,
            bg=self.c("panel_bg"),
            fg=self.c("text_fg"),
            padx=12,
            pady=12,
            bd=1,
            relief=tk.SOLID,
            font=("Arial", 12, "bold"),
            highlightbackground=self.c("border"),
            highlightcolor=self.c("border"),
        ), "section")

    def _label(self, parent, text: str, *, muted: bool = False) -> tk.Label:
        role = "muted_label" if muted else "label"
        return self._track(tk.Label(
            parent,
            text=text,
            bg=parent.cget("bg"),
            fg=self.c("muted_fg") if muted else self.c("text_fg"),
            anchor="w",
        ), role)

    def _entry(self, parent, variable: tk.StringVar, *, width: Optional[int] = None) -> tk.Entry:
        options = {}
        if width is not None:
            options["width"] = width
        entry = self._track(tk.Entry(
            parent,
            textvariable=variable,
            bg=self.c("text_bg"),
            fg=self.c("text_fg"),
            insertbackground=self.c("text_fg"),
            relief=tk.SOLID,
            bd=1,
            highlightthickness=1,
            highlightbackground=self.c("border"),
            highlightcolor=self.c("accent"),
            **options,
        ), "entry")
        entry.bind("<Return>", lambda _event: self.focus_get().tk_focusNext().focus_set() or "break")
        return entry

    def _button(self, parent, text: str, command, *, accent: bool = False) -> tk.Label:
        role = "accent_button" if accent else "button"
        button = self._track(tk.Label(
            parent,
            text=text,
            bg=self.c("accent") if accent else self.c("button_bg"),
            fg="#ffffff" if accent else self.c("text_fg"),
            activebackground=self.c("accent") if accent else self.c("button_active"),
            activeforeground="#ffffff" if accent else self.c("text_fg"),
            disabledforeground=self.c("disabled_fg"),
            relief=tk.RAISED,
            bd=1,
            padx=10,
            pady=5,
            cursor="hand2",
        ), role)
        button.normal_text = text
        button.is_accent = accent
        button.button_command = command
        button.bind("<Button-1>", self._invoke_label_button)
        button.bind("<Enter>", self._button_hover_enter)
        button.bind("<Leave>", self._button_hover_leave)
        return button

    def _invoke_label_button(self, event) -> None:
        widget = event.widget
        if str(widget.cget("state")) != tk.DISABLED:
            widget.button_command()

    def _button_hover_enter(self, event) -> None:
        widget = event.widget
        if str(widget.cget("state")) == tk.DISABLED:
            widget.configure(cursor="pirate")
        else:
            widget.configure(cursor="hand2")

    def _button_hover_leave(self, event) -> None:
        widget = event.widget
        widget.configure(text=widget.normal_text)
        self._apply_button_state()

    def _icon_button(self, parent, text: str, command) -> tk.Label:
        icon = self._track(tk.Label(
            parent,
            text=text,
            bg=parent.cget("bg"),
            fg=self.c("accent"),
            padx=2,
            cursor="hand2",
            font=("Arial", 14, "bold"),
        ), "icon_label")
        icon.bind("<Button-1>", lambda _event: command())
        return icon

    def _toggle_button(self, parent, variable: tk.BooleanVar, command) -> tk.Frame:
        frame = self._track(tk.Frame(parent, bg=parent.cget("bg"), cursor="hand2"), "toggle_frame")
        canvas = self._track(tk.Canvas(
            frame,
            width=48,
            height=26,
            bg=parent.cget("bg"),
            bd=0,
            highlightthickness=0,
            cursor="hand2",
        ), "toggle_canvas")
        label = self._track(tk.Label(
            frame,
            text="Light Mode",
            bg=parent.cget("bg"),
            fg=self.c("toggle_text"),
            padx=6,
            cursor="hand2",
            font=("Arial", 12),
        ), "toggle_label")
        canvas.pack(side=tk.LEFT)
        label.pack(side=tk.LEFT)
        for widget in (frame, canvas, label):
            widget.bind("<Button-1>", lambda _event: command())
        frame.toggle_canvas = canvas
        frame.toggle_label = label
        return frame

    def _choice_label(self, parent, text: str, variable: tk.BooleanVar, command) -> tk.Label:
        choice = self._track(tk.Label(
            parent,
            bg=parent.cget("bg"),
            fg=self.c("text_fg"),
            padx=2,
            cursor="hand2",
            font=("Arial", 13),
        ), "choice_label")
        choice.choice_text = text
        choice.choice_var = variable
        choice.bind("<Button-1>", lambda _event: command())
        return choice

    def _combo(self, parent, label: str, variable: tk.StringVar, values: tuple[str, ...], row: int, column: int) -> None:
        self._label(parent, label).grid(row=row, column=column, sticky="w", padx=(0, 6))
        dropdown = tk.OptionMenu(parent, variable, *values)
        dropdown.configure(
            bg=self.c("text_bg"),
            fg=self.c("text_fg"),
            activebackground=self.c("button_active"),
            activeforeground=self.c("text_fg"),
            relief=tk.SOLID,
            bd=1,
            highlightthickness=1,
            highlightbackground=self.c("border"),
            highlightcolor=self.c("accent"),
            anchor="w",
            width=14,
            cursor="hand2",
        )
        dropdown["menu"].configure(
            bg=self.c("text_bg"),
            fg=self.c("text_fg"),
            activebackground=self.c("button_active"),
            activeforeground=self.c("text_fg"),
        )
        self._track(dropdown, "option")
        self._track(dropdown["menu"], "menu")
        dropdown.grid(row=row, column=column + 1, sticky="ew", padx=(0, 12))

    def _text_box(
        self,
        parent,
        label: str,
        row: int,
        column: int,
        columnspan: int = 1,
        height: int = 8,
        footer_var: Optional[tk.StringVar] = None,
    ) -> tk.Text:
        self._label(parent, label).grid(row=row, column=column, columnspan=columnspan, sticky="w")
        frame = self._track(tk.Frame(parent, bg=self.c("panel_bg")), "panel_frame")
        frame.grid(row=row + 1, column=column, columnspan=columnspan, sticky="nsew", padx=(0, 8), pady=(4, 12))
        frame.columnconfigure(0, weight=1)
        frame.rowconfigure(0, weight=1)

        text = self._text_widget(frame, height=height)
        text.grid(row=0, column=0, sticky="nsew")
        if footer_var is not None:
            footer = self._label(frame, "", muted=True)
            footer.configure(textvariable=footer_var)
            footer.grid(row=1, column=0, sticky="w", pady=(4, 0))
        return text

    def _text_widget(self, parent, height: int) -> tk.Text:
        text = self._track(tk.Text(
            parent,
            height=height,
            wrap=tk.WORD,
            undo=True,
            bg=self.c("text_bg"),
            fg=self.c("text_fg"),
            insertbackground=self.c("text_fg"),
            selectbackground=self.c("select_bg"),
            selectforeground=self.c("text_fg"),
            relief=tk.SOLID,
            borderwidth=1,
            highlightthickness=1,
            highlightbackground=self.c("border"),
            highlightcolor=self.c("accent"),
            font=("Arial", 12),
        ), "text")
        text.bind("<Tab>", self._focus_next)
        text.bind("<Shift-Tab>", self._focus_previous)
        text.bind("<KeyRelease>", lambda _event: self._update_validation_state())
        text.bind("<FocusOut>", lambda _event: self._update_validation_state())
        return text

    def _track(self, widget, role: str):
        self._widgets_by_role.setdefault(role, []).append(widget)
        self._add_submit_bindtag(widget)
        return widget

    def _add_submit_bindtag(self, widget) -> None:
        try:
            bindtags = widget.bindtags()
        except tk.TclError:
            return
        if "SubmitShortcut" not in bindtags:
            widget.bindtags(("SubmitShortcut", *bindtags))

    def _focus_next(self, event) -> str:
        event.widget.tk_focusNext().focus_set()
        return "break"

    def _focus_previous(self, event) -> str:
        event.widget.tk_focusPrev().focus_set()
        return "break"

    def _bind_shortcuts(self) -> None:
        for sequence in (
            "<Command-Return>",
            "<Command-KeyPress-Return>",
            "<Control-Return>",
            "<Control-KeyPress-Return>",
            "<Command-KP_Enter>",
            "<Control-KP_Enter>",
        ):
            self.bind_class("SubmitShortcut", sequence, self._shortcut_generate)
            self.bind(sequence, self._shortcut_generate)

    def _wire_validation(self) -> None:
        for variable in (
            self.name_var,
            self.email_var,
            self.phone_var,
            self.location_var,
            self.output_dir_var,
            self.output_name_part_var,
            self.output_var,
            self.masters_var,
            self.bachelors_var,
        ):
            variable.trace_add("write", lambda *_args: self._update_validation_state())
        self.output_name_part_var.trace_add("write", lambda *_args: self._refresh_output_path())
        self.output_dir_var.trace_add("write", lambda *_args: self._refresh_output_path())

    def _shortcut_generate(self, _event) -> str:
        if self._is_form_valid():
            self._generate()
        return "break"

    def _text_value(self, text_widget: tk.Text) -> str:
        return text_widget.get("1.0", tk.END).strip()

    def _summary_word_count(self) -> int:
        return count_words(self._text_value(self.summary_text))

    def _validation_errors(self) -> list[str]:
        errors = []
        summary_words = self._summary_word_count()
        if not self._text_value(self.summary_text):
            errors.append("Summary is required")
        elif summary_words <= 40 or summary_words > 50:
            errors.append("Summary must be 41-50 words")
        if not self._text_value(self.skills_text):
            errors.append("Skills are required")
        if not self._text_value(self.experience_text):
            errors.append("Job Experience is required")
        if not self._text_value(self.top_skills_text):
            errors.append("Top 5 Skills are required")
        if not self.output_var.get().strip():
            errors.append("Output file is required")
        if self._matching_existing_output_files():
            errors.append("A matching DOCX or PDF file already exists")
        return errors

    def _is_form_valid(self) -> bool:
        return not self._validation_errors()

    def _update_validation_state(self) -> None:
        if not hasattr(self, "summary_text"):
            return
        summary_words = self._summary_word_count()
        summary_ok = 40 < summary_words <= 50
        self.summary_count_var.set(f"Summary: {summary_words} words (required: 41-50)")
        errors = self._validation_errors()
        if errors:
            self.validation_var.set("Generate disabled: " + "; ".join(errors))
            submit_state = tk.DISABLED
        else:
            self.validation_var.set("Ready to generate.")
            submit_state = tk.NORMAL
        for widget in self._submit_widgets:
            widget.configure(state=submit_state)
        self._refresh_custom_controls()
        self._apply_button_state()
        self._style_status_labels(summary_ok=summary_ok, form_ok=not errors)
        self._refresh_output_warning()

    def _style_status_labels(self, *, summary_ok: bool, form_ok: bool) -> None:
        for widget in self._widgets_by_role.get("muted_label", []):
            try:
                widget.configure(fg=self.c("muted_fg"))
            except tk.TclError:
                pass
        if hasattr(self, "validation_label"):
            self.validation_label.configure(fg=self.c("ok_fg") if form_ok else self.c("error_fg"))

    def _apply_button_state(self) -> None:
        for widget in self._widgets_by_role.get("accent_button", []):
            if str(widget.cget("state")) == tk.DISABLED:
                widget.configure(
                    bg=self.c("disabled_bg"),
                    fg=self.c("disabled_fg"),
                    activebackground=self.c("disabled_bg"),
                    activeforeground=self.c("disabled_fg"),
                    disabledforeground=self.c("disabled_fg"),
                    cursor="arrow",
                )
            else:
                widget.configure(
                    bg=self.c("accent"),
                    fg="#ffffff",
                    activebackground=self.c("accent"),
                    activeforeground="#ffffff",
                    disabledforeground=self.c("disabled_fg"),
                    cursor="hand2",
                )

    def _show_education_editor(self, education_key: str) -> None:
        self.education_editor_frame.grid()
        if education_key == "masters":
            self.education_editor_label.configure(text="Masters Education")
            self.bachelors_education_text.grid_remove()
            self.masters_education_text.grid()
        else:
            self.education_editor_label.configure(text="Bachelors Education")
            self.masters_education_text.grid_remove()
            self.bachelors_education_text.grid()

    def _toggle_masters(self) -> None:
        self.masters_var.set(not self.masters_var.get())
        self._update_validation_state()

    def _toggle_bachelors(self) -> None:
        self.bachelors_var.set(not self.bachelors_var.get())
        self._update_validation_state()

    def _refresh_custom_controls(self) -> None:
        if hasattr(self, "theme_toggle"):
            self._draw_theme_toggle()
        for widget in self._widgets_by_role.get("choice_label", []):
            is_checked = bool(widget.choice_var.get())
            marker = "☑" if is_checked else "☐"
            widget.configure(text=f"{marker} {widget.choice_text}")

    def _draw_theme_toggle(self) -> None:
        canvas = self.theme_toggle.toggle_canvas
        label = self.theme_toggle.toggle_label
        is_dark = self.dark_mode_var.get()
        canvas.delete("all")
        canvas.configure(bg=self.theme_toggle.cget("bg"))
        track = self.c("toggle_track") if is_dark else self.c("button_active")
        knob_x = 35 if is_dark else 13
        canvas.create_oval(1, 1, 25, 25, fill=track, outline=track)
        canvas.create_oval(23, 1, 47, 25, fill=track, outline=track)
        canvas.create_rectangle(13, 1, 35, 25, fill=track, outline=track)
        canvas.create_oval(knob_x - 10, 3, knob_x + 10, 23, fill=self.c("toggle_knob"), outline=self.c("toggle_knob"))
        label.configure(text="Dark Mode" if is_dark else "Light Mode", fg=self.c("toggle_text"))

    def _selected_education_entries(self) -> Tuple[str, ...]:
        entries = []
        if self.masters_var.get():
            entries.append(self._text_value(self.masters_education_text))
        if self.bachelors_var.get():
            entries.append(self._text_value(self.bachelors_education_text))
        return tuple(entry for entry in entries if entry.strip())

    def _toggle_theme(self) -> None:
        self.dark_mode_var.set(not self.dark_mode_var.get())
        self._configure_style()
        self._apply_theme()
        self._refresh_custom_controls()
        self._update_validation_state()

    def _apply_theme(self) -> None:
        role_options = {
            "app_frame": {"bg": self.c("app_bg")},
            "panel_frame": {"bg": self.c("panel_bg")},
            "section": {
                "bg": self.c("panel_bg"),
                "fg": self.c("text_fg"),
                "highlightbackground": self.c("border"),
                "highlightcolor": self.c("border"),
            },
            "label": {"fg": self.c("text_fg")},
            "muted_label": {"fg": self.c("muted_fg")},
            "entry": {
                "bg": self.c("text_bg"),
                "fg": self.c("text_fg"),
                "insertbackground": self.c("text_fg"),
                "highlightbackground": self.c("border"),
                "highlightcolor": self.c("accent"),
            },
            "text": {
                "bg": self.c("text_bg"),
                "fg": self.c("text_fg"),
                "insertbackground": self.c("text_fg"),
                "selectbackground": self.c("select_bg"),
                "selectforeground": self.c("text_fg"),
                "highlightbackground": self.c("border"),
                "highlightcolor": self.c("accent"),
            },
            "button": {
                "bg": self.c("button_bg"),
                "fg": self.c("text_fg"),
                "activebackground": self.c("button_active"),
                "activeforeground": self.c("text_fg"),
                "disabledforeground": self.c("disabled_fg"),
                "cursor": "hand2",
            },
            "choice_label": {"fg": self.c("text_fg")},
            "toggle_frame": {},
            "toggle_canvas": {"bg": self.c("panel_bg")},
            "toggle_label": {"fg": self.c("toggle_text")},
            "icon_label": {"fg": self.c("accent")},
            "option": {
                "bg": self.c("text_bg"),
                "fg": self.c("text_fg"),
                "activebackground": self.c("button_active"),
                "activeforeground": self.c("text_fg"),
                "highlightbackground": self.c("border"),
                "highlightcolor": self.c("accent"),
            },
            "menu": {
                "bg": self.c("text_bg"),
                "fg": self.c("text_fg"),
                "activebackground": self.c("button_active"),
                "activeforeground": self.c("text_fg"),
            },
        }
        parent_bg_roles = {
            "label",
            "muted_label",
            "toggle_frame",
            "toggle_canvas",
            "icon_label",
        }
        parent_active_bg_roles = {"choice_label", "toggle_label"}
        for role, widgets in self._widgets_by_role.items():
            for widget in widgets:
                options = dict(role_options.get(role, {}))
                if role in parent_bg_roles:
                    try:
                        options["bg"] = widget.master.cget("bg")
                    except tk.TclError:
                        pass
                if role in parent_active_bg_roles:
                    try:
                        options["bg"] = widget.master.cget("bg")
                        options["activebackground"] = widget.master.cget("bg")
                    except tk.TclError:
                        pass
                try:
                    widget.configure(**options)
                except tk.TclError:
                    pass

    def _output_directory(self) -> Path:
        return Path(self.output_dir_var.get().strip() or str(DEFAULT_OUTPUT_DIR)).expanduser()

    def _applied_output_directory(self) -> Path:
        return self._output_directory() / "#applied"

    def _ensure_applied_output_directory(self) -> None:
        self._applied_output_directory().mkdir(parents=True, exist_ok=True)

    def _output_name_part(self) -> str:
        return re.sub(r"[^A-Za-z0-9_-]+", "", self.output_name_part_var.get().strip())

    def _output_path(self) -> Path:
        return self._output_directory() / default_output_filename_for(self._output_name_part())

    def _refresh_output_path(self) -> None:
        if not hasattr(self, "output_var"):
            return
        if self._suppress_output_path_refresh:
            return
        new_path = str(self._output_path())
        if self.output_var.get() != new_path:
            self.output_var.set(new_path)

    def _matching_existing_output_files(self) -> list[Path]:
        try:
            self._ensure_applied_output_directory()
        except OSError:
            return []
        return find_matching_output_files(self._output_directory(), self._output_name_part())

    def _refresh_output_warning(self) -> None:
        if not hasattr(self, "output_warning_var"):
            return
        matches = self._matching_existing_output_files()
        if matches:
            output_dir = self._output_directory()
            shown = ", ".join(str(path.relative_to(output_dir)) if path.is_relative_to(output_dir) else path.name for path in matches[:3])
            extra = "" if len(matches) <= 3 else f" and {len(matches) - 3} more"
            self.output_warning_var.set(f"Warning: matching file may already exist: {shown}{extra}")
        else:
            self.output_warning_var.set("")
        if hasattr(self, "output_warning_label"):
            self.output_warning_label.configure(fg=self.c("error_fg") if matches else self.c("muted_fg"))
        if hasattr(self, "file_name_check_label"):
            if not self.output_name_part_var.get().strip():
                self.file_name_check_label.configure(text="", fg=self.c("ok_fg"))
            elif matches:
                self.file_name_check_label.configure(text="✕", fg=self.c("error_fg"))
            else:
                self.file_name_check_label.configure(text="✓", fg=self.c("ok_fg"))

    def _choose_output(self) -> None:
        selected = filedialog.askdirectory(
            title="Choose output folder",
            initialdir=str(self._output_directory()),
        )
        if selected:
            self.output_dir_var.set(selected)
            self._refresh_output_path()
            self._update_validation_state()

    def _generate(self) -> None:
        try:
            fmt = ResumeFormat(
                font_name=self.font_var.get(),
                name_size=int(self.name_size_var.get()),
                heading_size=int(self.heading_size_var.get()),
                body_size=int(self.body_size_var.get()),
            )
            content = ResumeContent(
                name=self.name_var.get().strip(),
                email=self.email_var.get().strip(),
                phone=self.phone_var.get().strip(),
                location=self.location_var.get().strip(),
                summary=self.summary_text.get("1.0", tk.END).strip(),
                skills=self.skills_text.get("1.0", tk.END).strip(),
                experience=self.experience_text.get("1.0", tk.END).strip(),
                certifications=self.certifications_text.get("1.0", tk.END).strip(),
                top_skills=self.top_skills_text.get("1.0", tk.END).strip(),
                education_entries=self._selected_education_entries(),
            )
            if not self._is_form_valid():
                messagebox.showerror("Cannot generate resume", self.validation_var.get())
                return
            output_path = Path(self.output_var.get()).expanduser()
            if output_path.suffix.lower() != ".docx":
                output_path = output_path.with_suffix(".docx")
            build_resume(content, fmt, output_path)
        except Exception as exc:
            messagebox.showerror("Could not generate resume", str(exc))
            return
        self._clear_text_inputs()
        messagebox.showinfo("Resume generated", f"Saved to:\n{output_path}")

    def _clear_text_inputs(self) -> None:
        for text_widget in (
            self.summary_text,
            self.skills_text,
            self.experience_text,
            self.certifications_text,
            self.top_skills_text,
        ):
            text_widget.delete("1.0", tk.END)
        self._suppress_output_path_refresh = True
        try:
            self.output_name_part_var.set("")
        finally:
            self._suppress_output_path_refresh = False
        self._update_validation_state()


if __name__ == "__main__":
    ResumeWriterApp().mainloop()
