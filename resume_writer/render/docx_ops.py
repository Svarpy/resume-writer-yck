"""Low-level python-docx helpers used by heuristic and template renderers."""

from __future__ import annotations

import re
from typing import Iterable, Optional

from formats_store import (
    ExperienceLayout,
    HeadingStyle,
    PageBorderStyle,
    SeparatorStyle,
    default_document_structure,
    normalize_line_spacing,
)

from resume_writer.render.models import ResumeFormat
from resume_writer.constants import DATE_RANGE_RE
from resume_writer.render.parse import (
    clean_lines,
    detect_known_company,
    fallback_role_date,
    looks_like_role_line,
    normalize_date_text,
    normalize_experience_bullet_text,
    split_dated_line,
    split_known_company_line,
)

# Populated by load_docx_dependencies()
Document = None
Inches = None
OxmlElement = None
Pt = None
WD_ALIGN_PARAGRAPH = None
qn = None


def load_docx_dependencies() -> None:
    global Document, Inches, OxmlElement, Pt, WD_ALIGN_PARAGRAPH, qn

    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt
    except ImportError as exc:
        raise RuntimeError("python-docx is required. Run: python3 -m pip install -r requirements.txt") from exc


def _alignment_from_name(name: str):
    mapping = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }
    return mapping.get((name or "left").strip().lower(), WD_ALIGN_PARAGRAPH.LEFT)


def set_run_font(run, font_name: str, size: int, bold: bool = False) -> None:
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    run.font.size = Pt(size)
    run.bold = bold


def format_paragraph(paragraph, before: float = 0, after: float = 3, alignment=None, line_spacing: float = 1.0) -> None:
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = normalize_line_spacing(line_spacing)
    if alignment is not None:
        paragraph.alignment = alignment


def add_border_rule(document: Document, fmt: Optional[ResumeFormat] = None) -> None:
    structure = (fmt.structure if fmt is not None else None) or default_document_structure()
    separator: SeparatorStyle = structure.separator
    if not separator.enabled:
        return
    paragraph = document.add_paragraph()
    format_paragraph(paragraph, before=0, after=3, line_spacing=fmt.line_spacing if fmt else 1.0)
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), separator.val or "single")
    bottom.set(qn("w:sz"), str(int(separator.sz)))
    bottom.set(qn("w:space"), str(int(separator.space)))
    bottom.set(qn("w:color"), separator.color or "AAAAAA")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def add_hyperlink(paragraph, text: str, url: str, fmt: ResumeFormat, bold: bool = False) -> None:
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")

    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), fmt.font_name)
    fonts.set(qn("w:hAnsi"), fmt.font_name)
    fonts.set(qn("w:eastAsia"), fmt.font_name)
    r_pr.append(fonts)

    size = OxmlElement("w:sz")
    size.set(qn("w:val"), str(fmt.body_size * 2))
    r_pr.append(size)

    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    r_pr.append(color)

    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)

    if bold:
        r_pr.append(OxmlElement("w:b"))

    new_run.append(r_pr)
    text_element = OxmlElement("w:t")
    text_element.text = text
    new_run.append(text_element)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_text_run(paragraph, text: str, fmt: ResumeFormat, bold: bool = False) -> None:
    run = paragraph.add_run(text)
    set_run_font(run, fmt.font_name, fmt.body_size, bold=bold)


def add_section_heading(document: Document, title: str, fmt: ResumeFormat) -> None:
    heading: HeadingStyle = fmt.structure.heading_style
    paragraph = document.add_paragraph()
    format_paragraph(
        paragraph,
        before=heading.space_before_pt,
        after=heading.space_after_pt,
        alignment=_alignment_from_name(heading.alignment),
        line_spacing=fmt.line_spacing,
    )
    run = paragraph.add_run(title)
    set_run_font(run, fmt.font_name, fmt.heading_size, bold=heading.bold)


def add_plain_paragraph(
    document: Document,
    text: str,
    fmt: ResumeFormat,
    alignment=None,
    before: float = 2,
    after: float = 4,
) -> None:
    paragraph = document.add_paragraph()
    if alignment is None:
        alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    format_paragraph(paragraph, before=before, after=after, alignment=alignment)
    add_text_run(paragraph, text, fmt)


def add_skill_line(document: Document, line: str, fmt: ResumeFormat) -> None:
    paragraph = document.add_paragraph()
    format_paragraph(paragraph, before=2.5, after=2.5, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    if ":" in line:
        label, rest = line.split(":", 1)
        add_text_run(paragraph, f"{label.strip()}: ", fmt, bold=True)
        add_text_run(paragraph, rest.strip(), fmt)
    else:
        add_text_run(paragraph, line, fmt)


def add_bullet(document: Document, text: str, fmt: ResumeFormat, *, normalize_dash: bool = False) -> None:
    if normalize_dash:
        text = normalize_experience_bullet_text(text)
    # Template packages (e.g. YcKResume) may omit the built-in List Bullet style.
    style_name = "List Bullet"
    try:
        _ = document.styles[style_name]
        paragraph = document.add_paragraph(style=style_name)
    except KeyError:
        paragraph = document.add_paragraph()
        marker = getattr(fmt.structure.experience, "bullet_marker", None) or "•"
        text = f"{marker}\t{text}"
    format_paragraph(paragraph, before=2, after=2, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    add_text_run(paragraph, text, fmt)


def set_cell_width(cell, width_inches: float) -> None:
    cell.width = Inches(width_inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(round(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_no_wrap(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    no_wrap = tc_pr.first_child_found_in("w:noWrap")
    if no_wrap is None:
        tc_pr.append(OxmlElement("w:noWrap"))


def set_table_grid(table, widths_inches: tuple[float, ...]) -> None:
    tbl_grid = table._tbl.tblGrid
    for grid_col in list(tbl_grid):
        tbl_grid.remove(grid_col)
    for width_inches in widths_inches:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(round(width_inches * 1440)))
        tbl_grid.append(grid_col)


def set_cell_margins(cell, *, top: int = 0, start: int = 0, bottom: int = 0, end: int = 0) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        margin = tc_mar.find(qn(f"w:{margin_name}"))
        if margin is None:
            margin = OxmlElement(f"w:{margin_name}")
            tc_mar.append(margin)
        margin.set(qn("w:w"), str(value))
        margin.set(qn("w:type"), "dxa")


def remove_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for border_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{border_name}"))
        if border is None:
            border = OxmlElement(f"w:{border_name}")
            borders.append(border)
        border.set(qn("w:val"), "nil")


def show_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for border_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{border_name}"))
        if border is None:
            border = OxmlElement(f"w:{border_name}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "6")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "FF0000")


def add_tabbed_line(document: Document, left_text: str, right_text: str, fmt: ResumeFormat) -> None:
    right_text = normalize_date_text(right_text) if DATE_RANGE_RE.search(right_text) else right_text.strip()
    experience: ExperienceLayout = fmt.structure.experience

    if experience.company_line_mode == "single_line":
        paragraph = document.add_paragraph()
        format_paragraph(paragraph, before=4, after=2, alignment=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=fmt.line_spacing)
        add_text_run(paragraph, left_text.strip(), fmt, bold=experience.company_bold)
        if right_text.strip():
            add_text_run(paragraph, f"  {right_text.strip()}", fmt, bold=experience.duration_bold)
        return

    if experience.company_line_mode == "tabbed_line":
        paragraph = document.add_paragraph()
        format_paragraph(paragraph, before=2, after=2, alignment=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=fmt.line_spacing)
        _ensure_right_tab(paragraph, int(getattr(experience, "tab_pos_twips", 9200) or 9200))
        add_text_run(paragraph, left_text.strip(), fmt, bold=experience.company_bold)
        if right_text.strip():
            add_text_run(paragraph, "\t" + right_text.strip(), fmt, bold=experience.duration_bold)
        return

    left_w = float(experience.left_width_in) or 4.45
    right_w = float(experience.right_width_in) or 2.55
    table = document.add_table(rows=1, cols=2)
    table.autofit = False
    set_table_grid(table, (left_w, right_w))
    remove_table_borders(table)

    left_cell, right_cell = table.rows[0].cells
    set_cell_width(left_cell, left_w)
    set_cell_width(right_cell, right_w)
    set_cell_margins(left_cell)
    set_cell_margins(right_cell)
    set_cell_no_wrap(right_cell)

    left_paragraph = left_cell.paragraphs[0]
    format_paragraph(left_paragraph, before=4, after=2, alignment=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=fmt.line_spacing)
    add_text_run(left_paragraph, left_text.strip(), fmt, bold=experience.company_bold)

    right_paragraph = right_cell.paragraphs[0]
    format_paragraph(
        right_paragraph,
        before=4,
        after=2,
        alignment=_alignment_from_name(experience.duration_align),
        line_spacing=fmt.line_spacing,
    )
    if right_text.strip():
        add_text_run(right_paragraph, right_text.strip(), fmt, bold=experience.duration_bold)


def _ensure_right_tab(paragraph, pos_twips: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    tabs = p_pr.find(qn("w:tabs"))
    if tabs is None:
        tabs = OxmlElement("w:tabs")
        p_pr.append(tabs)
    for existing in list(tabs):
        tabs.remove(existing)
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:pos"), str(int(pos_twips)))
    tabs.append(tab)


def strip_bullet_marker(line: str) -> str:
    return re.sub(r"^(?:[-*]|\u2022)\s+", "", line.strip())


def add_experience_section(document: Document, lines: Iterable[str], fmt: ResumeFormat) -> None:
    current_company = None
    role_index_by_company: dict[str, int] = {}

    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue

        clean_line = strip_bullet_marker(stripped)
        dated = split_dated_line(clean_line)
        if dated:
            left, right = dated
            add_tabbed_line(document, left, right, fmt)
            detected_company = detect_known_company(left)
            if detected_company:
                current_company = detected_company
                role_index_by_company.setdefault(current_company, 0)
            elif looks_like_role_line(left):
                role_index = role_index_by_company.get(current_company or "", 0)
                if fallback_role_date(current_company, role_index):
                    role_index_by_company[current_company] = role_index + 1
            continue

        known_company = split_known_company_line(clean_line)
        if known_company:
            left, right, company_key = known_company
            add_tabbed_line(document, left, right, fmt)
            current_company = company_key
            role_index_by_company.setdefault(current_company, 0)
            continue

        if looks_like_role_line(clean_line):
            role_index = role_index_by_company.get(current_company or "", 0)
            role_date = fallback_role_date(current_company, role_index)
            if role_date:
                add_tabbed_line(document, clean_line, role_date, fmt)
                role_index_by_company[current_company] = role_index + 1
                continue

        add_bullet(document, clean_line, fmt, normalize_dash=True)


def add_multiline_section(
    document: Document,
    lines: Iterable[str],
    fmt: ResumeFormat,
    *,
    skill_mode: bool = False,
    paragraph_alignment=None,
) -> None:
    if paragraph_alignment is None:
        paragraph_alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
        bullet_match = re.match(r"^(?:[-*]|\u2022)\s+(.*)$", stripped)
        if bullet_match:
            add_bullet(document, bullet_match.group(1).strip(), fmt)
        elif skill_mode:
            add_skill_line(document, stripped, fmt)
        elif "\t" in stripped:
            left, right = stripped.split("\t", 1)
            add_tabbed_line(document, left, right, fmt)
        elif split_dated_line(stripped):
            left, right = split_dated_line(stripped)
            add_tabbed_line(document, left, right, fmt)
        else:
            add_plain_paragraph(document, stripped, fmt, alignment=paragraph_alignment, before=2, after=3)


def add_education_entry(document: Document, education_text: str, fmt: ResumeFormat) -> None:
    for line in clean_lines(education_text):
        stripped = line.strip()
        if not stripped:
            continue
        dated = split_dated_line(stripped)
        if dated:
            left, right = dated
            add_tabbed_line(document, left, right, fmt)
        else:
            add_plain_paragraph(document, stripped, fmt, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=3)


def apply_document_defaults(document: Document, fmt: ResumeFormat) -> None:
    section = document.sections[0]
    section.page_width = Inches(fmt.page_width_in)
    section.page_height = Inches(fmt.page_height_in)
    section.top_margin = Inches(fmt.margin_top_in)
    section.right_margin = Inches(fmt.margin_right_in)
    section.bottom_margin = Inches(fmt.margin_bottom_in)
    section.left_margin = Inches(fmt.margin_left_in)
    section.header_distance = Inches(fmt.header_distance_in)
    section.footer_distance = Inches(fmt.footer_distance_in)

    normal = document.styles["Normal"]
    normal.font.name = fmt.font_name
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), fmt.font_name)
    normal.font.size = Pt(fmt.body_size)
    normal.paragraph_format.line_spacing = normalize_line_spacing(fmt.line_spacing)

    bullet = document.styles["List Bullet"]
    bullet.font.name = fmt.font_name
    bullet._element.rPr.rFonts.set(qn("w:eastAsia"), fmt.font_name)
    bullet.font.size = Pt(fmt.body_size)

    apply_page_border(document, fmt.structure.page_border)


def apply_page_border(document: Document, border: PageBorderStyle) -> None:
    if not border.enabled:
        return
    section = document.sections[0]
    sect_pr = section._sectPr
    pg_borders = sect_pr.find(qn("w:pgBorders"))
    if pg_borders is None:
        pg_borders = OxmlElement("w:pgBorders")
        sect_pr.append(pg_borders)
    for side in border.sides:
        el = pg_borders.find(qn(f"w:{side}"))
        if el is None:
            el = OxmlElement(f"w:{side}")
            pg_borders.append(el)
        el.set(qn("w:val"), border.val or "single")
        el.set(qn("w:sz"), str(int(border.sz)))
        el.set(qn("w:space"), str(int(border.space)))
        el.set(qn("w:color"), border.color or "000000")
