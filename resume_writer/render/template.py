"""Template-preserving resume generation.

When a format was uploaded from Word, the original ``.docx`` is stored beside
the format XML. Generate clones that package (keeping ``sectPr`` / styles /
numbering) and injects Writer field content into section regions.
"""

from __future__ import annotations

import re
from copy import deepcopy
from dataclasses import dataclass
from pathlib import Path
from typing import TYPE_CHECKING, Callable, Optional

if TYPE_CHECKING:
    from resume_writer.render.models import ResumeContent, ResumeFormat


@dataclass
class TemplatePrototypes:
    separator_element: Optional[object] = None
    heading_element: Optional[object] = None
    dated_line_element: Optional[object] = None
    bullet_element: Optional[object] = None
    experience_table_element: Optional[object] = None


def build_resume_from_template(content: "ResumeContent", fmt: "ResumeFormat", output_path: Path) -> Path:
    """Clone the stored template ``.docx`` and replace body content."""
    from docx import Document
    from docx.oxml.ns import qn

    from resume_writer.constants import AUTHOR_NAME
    from resume_writer.render.docx_ops import load_docx_dependencies
    from resume_writer.render.heuristic import _populate_resume_body_for_template
    from resume_writer.render.parse import split_keywords

    load_docx_dependencies()
    template_path = Path(fmt.template_path)
    if not template_path.is_file():
        raise FileNotFoundError(f"Template not found: {template_path}")

    document = Document(str(template_path))
    structure = fmt.structure
    prototypes = capture_template_prototypes(document, structure)

    body = document.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)

    _populate_resume_body_for_template(document, content, fmt, prototypes=prototypes)

    core = document.core_properties
    core.author = content.name or AUTHOR_NAME
    core.subject = "my resume"
    core.keywords = ", ".join(split_keywords(content.top_skills))
    core.title = f"{content.name or AUTHOR_NAME} Resume"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path


def capture_template_prototypes(document, structure) -> TemplatePrototypes:
    protos = TemplatePrototypes()
    body_blocks = list(document.element.body)

    if structure is not None and 0 <= getattr(structure, "separator_block_index", -1) < len(body_blocks):
        protos.separator_element = deepcopy(body_blocks[structure.separator_block_index])

    if structure is not None:
        for section in structure.sections:
            idx = getattr(section, "template_block_index", -1)
            if 0 <= idx < len(body_blocks):
                protos.heading_element = deepcopy(body_blocks[idx])
                break
        tbl_idx = getattr(structure, "experience_table_index", -1)
        if 0 <= tbl_idx < len(body_blocks):
            protos.experience_table_element = deepcopy(body_blocks[tbl_idx])

    for paragraph in document.paragraphs:
        text = paragraph.text
        if protos.separator_element is None and _paragraph_is_separator(paragraph):
            protos.separator_element = deepcopy(paragraph._p)
        if protos.dated_line_element is None and "\t" in text and re.search(r"\d{4}", text):
            protos.dated_line_element = deepcopy(paragraph._p)
        if protos.bullet_element is None and re.match(r"^[•▸●○]\s*", text.strip()):
            protos.bullet_element = deepcopy(paragraph._p)

    if protos.heading_element is None and structure is not None:
        for section in structure.sections:
            for paragraph in document.paragraphs:
                if paragraph.text.strip() == section.heading:
                    protos.heading_element = deepcopy(paragraph._p)
                    break
            if protos.heading_element is not None:
                break

    return protos


def _paragraph_is_separator(paragraph) -> bool:
    from docx.oxml.ns import qn

    if paragraph.text.strip():
        return False
    p_pr = paragraph._p.pPr
    if p_pr is None:
        return False
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        return False
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        return False
    val = bottom.get(qn("w:val")) or ""
    return val not in {"", "nil", "none"}


def insert_before_sectpr(document, element) -> None:
    from docx.oxml.ns import qn

    body = document.element.body
    sect_pr = body.find(qn("w:sectPr"))
    if sect_pr is not None:
        sect_pr.addprevious(element)
    else:
        body.append(element)


def clone_separator(document, prototypes: Optional[TemplatePrototypes], fmt: "ResumeFormat") -> None:
    from resume_writer.render.docx_ops import add_border_rule

    if prototypes and prototypes.separator_element is not None:
        insert_before_sectpr(document, deepcopy(prototypes.separator_element))
        return
    add_border_rule(document, fmt)


def clone_heading(document, text: str, prototypes: Optional[TemplatePrototypes], fmt: "ResumeFormat") -> None:
    from resume_writer.render.docx_ops import add_section_heading

    if prototypes and prototypes.heading_element is not None:
        element = deepcopy(prototypes.heading_element)
        set_paragraph_element_text(element, text, bold=True)
        insert_before_sectpr(document, element)
        return
    add_section_heading(document, text, fmt)


def set_paragraph_element_text(p_element, text: str, *, bold: Optional[bool] = None) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    first_r_pr = None
    for child in list(p_element):
        tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
        if tag == "r":
            r_pr = child.find(qn("w:rPr"))
            if r_pr is not None and first_r_pr is None:
                first_r_pr = deepcopy(r_pr)
            p_element.remove(child)
        elif tag == "hyperlink":
            p_element.remove(child)

    run = OxmlElement("w:r")
    if first_r_pr is not None:
        if bold is True and first_r_pr.find(qn("w:b")) is None:
            first_r_pr.append(OxmlElement("w:b"))
        run.append(first_r_pr)
    elif bold:
        r_pr = OxmlElement("w:rPr")
        r_pr.append(OxmlElement("w:b"))
        run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.set(qn("xml:space"), "preserve")
    text_el.text = text
    run.append(text_el)
    p_element.append(run)
