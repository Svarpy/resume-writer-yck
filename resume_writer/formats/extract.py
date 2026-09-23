"""Extract resume format structure from uploaded Word documents."""

from __future__ import annotations

import re
import shutil
from pathlib import Path
from typing import Optional

from resume_writer.formats.models import (
    DEFAULT_FORMAT_VALUES,
    DocumentStructure,
    ExperienceLayout,
    FormatError,
    FormatValidationError,
    HeadingStyle,
    PageBorderStyle,
    ResumeFormatSpec,
    SectionRule,
    SeparatorStyle,
    _match_section_key,
    canonicalize_section_key,
    default_document_structure,
    normalize_line_spacing,
)
from resume_writer.formats.store import (
    _format_path,
    _new_format_id,
    _safe_format_id,
    _template_filename,
    create_format,
    delete_format,
    format_spec_to_xml,
    format_template_path,
    get_format,
)
from resume_writer.paths import sanitize_username, user_formats_dir

def _paragraph_has_bottom_border(paragraph) -> bool:
    try:
        from docx.oxml.ns import qn
    except ImportError:
        return False
    p_pr = paragraph._p.pPr
    if p_pr is None:
        return False
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        return False
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        return False
    val = bottom.get(qn("w:val")) or ""
    return val not in {"", "nil", "none"}


def _read_bottom_border(paragraph) -> Optional[SeparatorStyle]:
    try:
        from docx.oxml.ns import qn
    except ImportError:
        return None
    p_pr = paragraph._p.pPr
    if p_pr is None:
        return None
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        return None
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        return None
    val = bottom.get(qn("w:val")) or "single"
    if val in {"nil", "none"}:
        return None
    try:
        sz = int(bottom.get(qn("w:sz")) or "2")
    except ValueError:
        sz = 2
    try:
        space = int(bottom.get(qn("w:space")) or "1")
    except ValueError:
        space = 1
    color = bottom.get(qn("w:color")) or "AAAAAA"
    return SeparatorStyle(enabled=True, val=val, sz=sz, space=space, color=color)


def _looks_like_section_heading(paragraph, *, body_size: int) -> bool:
    text = paragraph.text.strip()
    if not text or len(text) > 48 or "\n" in text:
        return False
    if _match_section_key(text):
        return True
    # Short bold / larger / ALL-CAPS lines are likely headings.
    # Require ALL-CAPS for non-alias matches so name/body lines (e.g. "Jane Doe")
    # are not stored as synthetic section keys that confuse generation.
    runs = [r for r in paragraph.runs if r.text and r.text.strip()]
    if not runs:
        return False
    boldish = any(bool(r.bold) for r in runs)
    sizes = []
    for run in runs:
        if run.font.size is not None:
            sizes.append(int(round(run.font.size.pt)))
    larger = any(s >= body_size + 1 for s in sizes)
    mostly_upper = text == text.upper() and any(c.isalpha() for c in text)
    word_count = len(text.split())
    return word_count <= 6 and mostly_upper and (boldish or larger)


def _extract_page_border(section) -> PageBorderStyle:
    try:
        from docx.oxml.ns import qn
    except ImportError:
        return PageBorderStyle(enabled=False)

    sect_pr = section._sectPr
    if sect_pr is None:
        return PageBorderStyle(enabled=False)
    pg_borders = sect_pr.find(qn("w:pgBorders"))
    if pg_borders is None:
        return PageBorderStyle(enabled=False)

    sides: list[str] = []
    sample = None
    for side in ("top", "left", "bottom", "right"):
        el = pg_borders.find(qn(f"w:{side}"))
        if el is None:
            continue
        val = el.get(qn("w:val")) or ""
        if val in {"", "nil", "none"}:
            continue
        sides.append(side)
        if sample is None:
            sample = el
    if not sides or sample is None:
        return PageBorderStyle(enabled=False)
    try:
        sz = int(sample.get(qn("w:sz")) or "4")
    except ValueError:
        sz = 4
    try:
        space = int(sample.get(qn("w:space")) or "24")
    except ValueError:
        space = 24
    return PageBorderStyle(
        enabled=True,
        val=sample.get(qn("w:val")) or "single",
        sz=sz,
        space=space,
        color=sample.get(qn("w:color")) or "000000",
        sides=tuple(sides),
    )


def _extract_experience_layout(document) -> ExperienceLayout:
    layout = ExperienceLayout()
    try:
        from docx.oxml.ns import qn
    except ImportError:
        return layout

    for table in document.tables:
        if len(table.columns) != 2 or not table.rows:
            continue
        row = table.rows[0]
        if len(row.cells) < 2:
            continue
        left_text = row.cells[0].text.strip()
        right_text = row.cells[1].text.strip()
        if not left_text:
            continue
        # Prefer tables that look like company | duration.
        looks_dated = bool(re.search(r"\d{4}|present|current", right_text, re.IGNORECASE))
        if not looks_dated and len(right_text) > 40:
            continue

        widths: list[float] = []
        for cell in row.cells[:2]:
            tc_pr = cell._tc.tcPr
            width_in = None
            if tc_pr is not None:
                tc_w = tc_pr.first_child_found_in("w:tcW")
                if tc_w is not None:
                    try:
                        width_in = int(tc_w.get(qn("w:w")) or "0") / 1440.0
                    except ValueError:
                        width_in = None
            widths.append(width_in if width_in and width_in > 0 else 0.0)

        if widths[0] > 0 and widths[1] > 0:
            left_w, right_w = round(widths[0], 2), round(widths[1], 2)
            # Some templates report full-page cell widths for both columns; clamp
            # to a usable two-column split so generate never overflows page grid.
            total = left_w + right_w
            if total > 7.5:
                scale = 7.0 / total
                left_w = round(left_w * scale, 2)
                right_w = round(right_w * scale, 2)
            layout.left_width_in = left_w
            layout.right_width_in = right_w
        layout.company_line_mode = "two_column_table"

        # Detect bold on left/right runs.
        left_bold = any(bool(r.bold) for p in row.cells[0].paragraphs for r in p.runs)
        right_bold = any(bool(r.bold) for p in row.cells[1].paragraphs for r in p.runs)
        layout.company_bold = left_bold if row.cells[0].paragraphs else True
        layout.duration_bold = right_bold if row.cells[1].paragraphs else True

        right_align = "right"
        for paragraph in row.cells[1].paragraphs:
            align = paragraph.alignment
            if align is not None:
                # WD_ALIGN_PARAGRAPH.RIGHT == 2
                right_align = "right" if int(align) == 2 else "left" if int(align) == 0 else "center"
                break
        layout.duration_align = right_align
        return layout

    # YcKResume / YcKResumeFTR2: company|location \t duration with a right tab stop.
    for paragraph in document.paragraphs:
        raw = paragraph.text
        if "\t" not in raw:
            continue
        if not re.search(r"\d{4}|present|current", raw, re.IGNORECASE):
            continue
        p_pr = paragraph._p.pPr
        tab_pos = None
        if p_pr is not None:
            tabs_el = p_pr.find(qn("w:tabs"))
            if tabs_el is not None:
                for tab in tabs_el:
                    if (tab.get(qn("w:val")) or "") == "right":
                        try:
                            tab_pos = int(tab.get(qn("w:pos")) or "0")
                        except ValueError:
                            tab_pos = None
                        if tab_pos:
                            break
        layout.company_line_mode = "tabbed_line"
        layout.duration_align = "right"
        if tab_pos and tab_pos > 0:
            layout.tab_pos_twips = tab_pos
        layout.company_bold = any(bool(r.bold) for r in paragraph.runs) or True
        for para in document.paragraphs:
            t = para.text.strip()
            m = re.match(r"^([•▸●○\-–—*])\s*\t?\s*\S", t)
            if m:
                layout.bullet_marker = m.group(1)
                break
        return layout

    layout.company_line_mode = "two_column_table"
    return layout


def _extract_structure_from_document(document, *, body_size: int) -> DocumentStructure:
    structure = default_document_structure()
    sections: list[SectionRule] = []
    seen_keys: set[str] = set()
    separator: Optional[SeparatorStyle] = None
    pending_separator = False
    separator_block_index = -1
    order = 0
    heading_entries: list[tuple[SectionRule, int]] = []

    body_blocks = list(document.element.body)
    for block_index, block in enumerate(body_blocks):
        tag = block.tag.split("}")[-1] if "}" in block.tag else block.tag
        if tag == "tbl" and structure.experience_table_index < 0:
            for table in document.tables:
                if table._tbl is block and len(table.columns) == 2:
                    structure.experience_table_index = block_index
                    break
            continue
        if tag != "p":
            continue

        paragraph = None
        for candidate in document.paragraphs:
            if candidate._p is block:
                paragraph = candidate
                break
        if paragraph is None:
            continue

        if _paragraph_has_bottom_border(paragraph) and not paragraph.text.strip():
            pending_separator = True
            border = _read_bottom_border(paragraph)
            if border is not None:
                separator = border
                if separator_block_index < 0:
                    separator_block_index = block_index
            continue

        if not _looks_like_section_heading(paragraph, body_size=body_size):
            if _paragraph_has_bottom_border(paragraph):
                border = _read_bottom_border(paragraph)
                if border is not None:
                    separator = border
                    if separator_block_index < 0:
                        separator_block_index = block_index
            continue

        heading_text = paragraph.text.strip()
        key = _match_section_key(heading_text)
        if key is None:
            key = re.sub(r"[^a-z0-9]+", "_", heading_text.casefold()).strip("_") or f"section_{order}"
        key = canonicalize_section_key(key, heading_text)
        if key in seen_keys:
            continue
        seen_keys.add(key)
        rule = SectionRule(
            key=key,
            heading=heading_text,
            separator_before=pending_separator or _paragraph_has_bottom_border(paragraph),
            order=order,
            template_block_index=block_index,
            content_start_index=block_index + 1,
        )
        sections.append(rule)
        heading_entries.append((rule, block_index))
        order += 1
        pending_separator = False

        if order == 1:
            pf = paragraph.paragraph_format
            before = pf.space_before.pt if pf.space_before is not None else structure.heading_style.space_before_pt
            after = pf.space_after.pt if pf.space_after is not None else structure.heading_style.space_after_pt
            run_list = list(paragraph.runs)
            bold = any(bool(r.bold) for r in run_list) if run_list else True
            alignment = "left"
            if paragraph.alignment is not None:
                align_map = {0: "left", 1: "center", 2: "right", 3: "justify"}
                alignment = align_map.get(int(paragraph.alignment), "left")
            structure.heading_style = HeadingStyle(
                bold=bool(bold),
                space_before_pt=float(before),
                space_after_pt=float(after),
                alignment=alignment,
            )

    for idx, (rule, _heading_idx) in enumerate(heading_entries):
        end_idx = len(body_blocks) - 1
        next_boundary = heading_entries[idx + 1][1] if idx + 1 < len(heading_entries) else end_idx
        rule.content_end_index = next_boundary

    template_had_separators = separator is not None or any(s.separator_before for s in sections)
    if sections:
        for default_sec in default_document_structure().sections:
            if default_sec.key not in seen_keys:
                sections.append(
                    SectionRule(
                        key=default_sec.key,
                        heading=default_sec.heading,
                        separator_before=bool(template_had_separators),
                        order=len(sections),
                    )
                )
        structure.sections = sections

    structure.separator_block_index = separator_block_index
    if separator is not None:
        structure.separator = separator
    elif template_had_separators:
        structure.separator = SeparatorStyle(enabled=True)
    else:
        structure.separator = SeparatorStyle(enabled=False)

    if len(document.paragraphs) >= 2:
        contact_text = document.paragraphs[1].text
        if " · " in contact_text or "·" in contact_text:
            structure.contact_separator = "   ·   "
        elif "|" in contact_text:
            structure.contact_separator = "  |  "

    structure.experience = _extract_experience_layout(document)
    if document.sections:
        structure.page_border = _extract_page_border(document.sections[0])
    return structure


def _effective_run_font_name(run) -> Optional[str]:
    """Return the concrete font name on a run, including theme/ascii fallbacks."""
    name = run.font.name
    if name:
        return name
    try:
        from docx.oxml.ns import qn

        r_pr = run._element.rPr
        if r_pr is None:
            return None
        r_fonts = r_pr.rFonts
        if r_fonts is None:
            return None
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            value = r_fonts.get(qn(attr))
            if value:
                return value
    except Exception:
        return None
    return None


def extract_format_from_docx(docx_path: str | Path, *, name: str = "") -> ResumeFormatSpec:
    """Inspect a .docx and build a ResumeFormatSpec (not yet saved).

    Extracts page size/margins, typography hints, and structural rules
    (section headings, separators, experience layout, page borders).
    Missing values fall back to defaults.
    """
    path = Path(docx_path).expanduser()
    if not path.is_file():
        raise FormatValidationError(f"File not found: {path}")
    if path.suffix.casefold() != ".docx":
        raise FormatValidationError("Only .docx files are supported for format upload")

    try:
        from docx import Document
    except ImportError as exc:
        raise FormatValidationError(
            "python-docx is required. Run: python3 -m pip install -r requirements.txt"
        ) from exc

    document = Document(str(path))
    section = document.sections[0]

    def inches(value: Any, default: float) -> float:
        try:
            return round(float(value) / 914400, 4)  # EMUs → inches
        except Exception:
            return default

    font_name: Optional[str] = None
    body_size = DEFAULT_FORMAT_VALUES["body_size"]
    name_size = DEFAULT_FORMAT_VALUES["name_size"]
    heading_size = DEFAULT_FORMAT_VALUES["heading_size"]

    try:
        normal = document.styles["Normal"]
        if normal.font.name:
            font_name = normal.font.name
        if normal.font.size is not None:
            body_size = int(round(normal.font.size.pt))
    except Exception:
        pass

    # Scan early runs for fonts and sizes. Real run fonts must win over the
    # default / Normal fallback (previously `font_name = font_name or run.font.name`
    # left Arial stuck once the default was assigned).
    sizes: list[int] = []
    bold_sizes: list[int] = []
    run_fonts: list[str] = []
    for paragraph in document.paragraphs[:80]:
        for run in paragraph.runs:
            run_font = _effective_run_font_name(run)
            if run_font:
                run_fonts.append(run_font)
                font_name = run_font
            if run.font.size is not None:
                pt = int(round(run.font.size.pt))
                sizes.append(pt)
                if run.bold:
                    bold_sizes.append(pt)

    if run_fonts:
        # Prefer the most common explicit run font.
        font_name = max(set(run_fonts), key=run_fonts.count)

    if bold_sizes:
        name_size = max(bold_sizes)
    if sizes:
        above_body = [s for s in sizes if s > body_size]
        if above_body:
            heading_size = sorted(above_body)[len(above_body) // 2]
        if not bold_sizes:
            name_size = max(sizes)

    structure = _extract_structure_from_document(document, body_size=body_size)

    # Infer line spacing from Normal style when available.
    line_spacing = DEFAULT_FORMAT_VALUES["line_spacing"]
    try:
        normal = document.styles["Normal"]
        if normal.paragraph_format.line_spacing is not None:
            line_spacing = normalize_line_spacing(normal.paragraph_format.line_spacing)
    except Exception:
        pass

    display_name = name.strip() or path.stem
    return ResumeFormatSpec(
        id=_new_format_id(display_name),
        name=display_name,
        font_name=font_name or DEFAULT_FORMAT_VALUES["font_name"],
        name_size=name_size,
        heading_size=heading_size,
        body_size=body_size,
        page_width_in=inches(section.page_width, DEFAULT_FORMAT_VALUES["page_width_in"]),
        page_height_in=inches(section.page_height, DEFAULT_FORMAT_VALUES["page_height_in"]),
        margin_top_in=inches(section.top_margin, DEFAULT_FORMAT_VALUES["margin_top_in"]),
        margin_right_in=inches(section.right_margin, DEFAULT_FORMAT_VALUES["margin_right_in"]),
        margin_bottom_in=inches(section.bottom_margin, DEFAULT_FORMAT_VALUES["margin_bottom_in"]),
        margin_left_in=inches(section.left_margin, DEFAULT_FORMAT_VALUES["margin_left_in"]),
        header_distance_in=inches(section.header_distance, DEFAULT_FORMAT_VALUES["header_distance_in"]),
        footer_distance_in=inches(section.footer_distance, DEFAULT_FORMAT_VALUES["footer_distance_in"]),
        line_spacing=line_spacing,
        structure=structure,
        protected=False,
        source="docx",
    )


def add_format_from_docx(
    username: str,
    docx_path: str | Path,
    name: str,
    *,
    set_as_primary: bool = False,
) -> ResumeFormatSpec:
    """Add Format flow: extract structure, store the source ``.docx``, save under ``name``.

    The uploaded Word package is copied beside the format XML so generate can
    clone it (template-preserving path) instead of approximating layout.
    """
    key = sanitize_username(username)
    source_path = Path(docx_path).expanduser()
    extracted = extract_format_from_docx(source_path, name=name)
    spec = create_format(
        key,
        name,
        xml_text=format_spec_to_xml(extracted),
        source="docx",
        set_as_primary=set_as_primary,
    )

    template_name = _template_filename(spec.id)
    dest = user_formats_dir(key) / template_name
    try:
        shutil.copy2(source_path, dest)
    except OSError as exc:
        try:
            delete_format(key, spec.id)
        except FormatError:
            pass
        raise FormatValidationError(f"Could not store template document: {exc}") from exc

    spec.template_file = template_name
    _format_path(key, spec.id).write_text(format_spec_to_xml(spec), encoding="utf-8")
    return spec

