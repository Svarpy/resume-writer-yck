"""Smoke checks for Sprint 1 auth + formats store (no UI)."""

from __future__ import annotations

import os
import tempfile
import unittest
from pathlib import Path

# Isolate all tests from the developer's real ~/.ResumeWriter data.
_TEST_ROOT = tempfile.mkdtemp(prefix="rw-sprint1-")
os.environ["RESUME_WRITER_DATA_DIR"] = _TEST_ROOT

from formats_store import (  # noqa: E402
    DEFAULT_FORMAT_ID,
    FormatProtectedError,
    add_format_from_docx,
    create_format,
    default_format_spec,
    default_format_xml,
    delete_format,
    extract_format_from_docx,
    format_spec_from_xml,
    get_format,
    get_primary_format,
    list_formats,
    set_primary_format,
    update_format,
)
from user_auth import (  # noqa: E402
    InvalidCredentialsError,
    PasswordValidationError,
    UserExistsError,
    change_password,
    clear_session,
    get_current_user,
    get_settings,
    signin,
    signup,
    update_profile,
)


class AuthTests(unittest.TestCase):
    def test_signup_signin_change_password_and_settings(self) -> None:
        profile = signup("DemoUser", "Secret123", display_name="Demo", dark_mode=True)
        self.assertEqual(profile.username, "demouser")
        self.assertEqual(get_current_user(), "demouser")
        self.assertNotIn("password", profile.public_dict())
        self.assertTrue(profile.password_hash.startswith("$2"))

        clear_session()
        self.assertIsNone(get_current_user())
        signed = signin("DemoUser", "Secret123")
        self.assertEqual(signed.username, "demouser")

        with self.assertRaises(UserExistsError):
            signup("demouser", "Other123")

        with self.assertRaises(InvalidCredentialsError):
            signin("demouser", "wrongpass1")

        change_password("demouser", "Secret123", "Newer456")
        with self.assertRaises(InvalidCredentialsError):
            signin("demouser", "Secret123")
        signin("demouser", "Newer456")

        update_profile("demouser", dark_mode=False, primary_format_id="default")
        settings = get_settings("demouser")
        self.assertFalse(settings["dark_mode"])
        self.assertEqual(settings["primary_format_id"], "default")

    def test_password_validation(self) -> None:
        with self.assertRaises(PasswordValidationError):
            signup("shortuser", "abc")
        with self.assertRaises(PasswordValidationError):
            signup("noletter1", "12345678")


class FormatTests(unittest.TestCase):
    def setUp(self) -> None:
        if get_current_user() != "formatuser":
            try:
                signup("formatuser", "Format123")
            except UserExistsError:
                signin("formatuser", "Format123")

    def test_default_is_always_listed_and_protected(self) -> None:
        formats = list_formats("formatuser")
        self.assertEqual(formats[0].id, DEFAULT_FORMAT_ID)
        self.assertTrue(formats[0].protected)
        xml = default_format_xml()
        self.assertIn("<resumeFormat", xml)
        with self.assertRaises(FormatProtectedError):
            update_format("formatuser", DEFAULT_FORMAT_ID, name="Nope")
        with self.assertRaises(FormatProtectedError):
            delete_format("formatuser", DEFAULT_FORMAT_ID)

    def test_crud_and_primary(self) -> None:
        created = create_format(
            "formatuser",
            "Compact",
            values={"font_name": "Calibri", "body_size": 10},
            set_as_primary=True,
        )
        self.assertEqual(get_primary_format("formatuser").id, created.id)
        self.assertEqual(get_format("formatuser", created.id).font_name, "Calibri")

        updated = update_format("formatuser", created.id, values={"heading_size": 14})
        self.assertEqual(updated.heading_size, 14)

        set_primary_format("formatuser", DEFAULT_FORMAT_ID)
        delete_format("formatuser", created.id)
        ids = {f.id for f in list_formats("formatuser")}
        self.assertNotIn(created.id, ids)
        self.assertEqual(get_primary_format("formatuser").id, DEFAULT_FORMAT_ID)

    def test_xml_roundtrip(self) -> None:
        spec = default_format_spec()
        restored = format_spec_from_xml(default_format_xml())
        self.assertEqual(restored.font_name, spec.font_name)
        self.assertEqual(restored.name_size, spec.name_size)
        self.assertEqual(restored.margin_top_in, spec.margin_top_in)


class DocxExtractTests(unittest.TestCase):
    def test_extract_from_generated_docx(self) -> None:
        # Build a tiny docx with python-docx mirroring writer defaults.
        from docx import Document
        from docx.shared import Inches, Pt

        doc = Document()
        section = doc.sections[0]
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        p = doc.add_paragraph()
        run = p.add_run("Sample Name")
        run.bold = True
        run.font.size = Pt(16)
        run.font.name = "Arial"

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "sample.docx"
            doc.save(path)
            spec = extract_format_from_docx(path, name="FromDocx")
            self.assertEqual(spec.source, "docx")
            self.assertAlmostEqual(spec.page_width_in, 8.5, places=2)
            self.assertEqual(spec.name_size, 16)

            signup_name = "docxuser"
            try:
                signup(signup_name, "DocxUser1")
            except UserExistsError:
                signin(signup_name, "DocxUser1")
            saved = add_format_from_docx(signup_name, path, "Uploaded Layout")
            self.assertEqual(saved.name, "Uploaded Layout")
            self.assertFalse(saved.protected)


if __name__ == "__main__":
    unittest.main()
