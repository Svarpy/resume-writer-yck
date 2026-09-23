"""Template-preserving format extract → primary → generate tests.

Uses the PM reference resumes under tests/fixtures/formats/ as ground truth:
- YcKResume.docx — 0.5" margins, Title Case headings, border separators, tabbed experience
- YcKResumeFTR2.docx — compact margins, ALL-CAPS headings, no separators, tabbed experience
"""

from __future__ import annotations


# Ensure repo root is importable when tests live under tests/.
import sys
from pathlib import Path as _Path
_ROOT = _Path(__file__).resolve().parents[1]
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

import os
import tempfile
import unittest
from pathlib import Path

_TEST_ROOT = tempfile.mkdtemp(prefix="rw-template-")
os.environ["RESUME_WRITER_DATA_DIR"] = _TEST_ROOT

from formats_store import (  # noqa: E402
    DEFAULT_FORMAT_ID,
    add_format_from_docx,
    extract_format_from_docx,
    get_primary_format,
    resolve_format_template_path,
    set_primary_format,
)
from user_auth import UserExistsError, signin, signup  # noqa: E402

FIXTURES = Path(__file__).resolve().parent / "fixtures" / "formats"
YCK_RESUME = FIXTURES / "YcKResume.docx"
YCK_FTR2 = FIXTURES / "YcKResumeFTR2.docx"


def _sample_content():
    from resume_writer_app import ResumeContent

    return ResumeContent(
        name="Test User",
        email="t@example.com",
        phone="+1 555 0100",
        location="Remote",
        summary="Word " * 45,
        skills="Languages: Python, JavaScript\nCloud: AWS, Docker",
        experience=(
            "Acme Corp | Remote\tJan 2020 – Present\n"
            "Software Engineer\n"
            "- Built things\n"
            "- Shipped features"
        ),
        certifications="",
        top_skills="Python, Word",
        education_entries=("State University\tMay 2018 – May 2022",),
    )


def _ensure_user(name: str, password: str) -> None:
    try:
        signup(name, password)
    except UserExistsError:
        signin(name, password)


class ReferenceExtractTests(unittest.TestCase):
    def test_yck_resume_extracts_margins_headings_separator_and_tabs(self) -> None:
        self.assertTrue(YCK_RESUME.is_file())
        spec = extract_format_from_docx(YCK_RESUME, name="YcKResume")
        self.assertAlmostEqual(spec.margin_top_in, 0.5, places=2)
        self.assertAlmostEqual(spec.margin_left_in, 0.5, places=2)
        self.assertNotAlmostEqual(spec.margin_top_in, 0.75, places=2)

        headings = {s.key: s.heading for s in spec.structure.sections}
        self.assertEqual(headings.get("summary"), "Summary")
        self.assertEqual(headings.get("skills"), "Skills")
        self.assertEqual(headings.get("experience"), "Professional Experience")
        self.assertEqual(headings.get("education"), "Education")

        self.assertTrue(spec.structure.separator.enabled)
        self.assertEqual(spec.structure.separator.color, "000000")
        self.assertEqual(spec.structure.experience.company_line_mode, "tabbed_line")
        self.assertEqual(spec.structure.experience.tab_pos_twips, 9200)
        self.assertGreaterEqual(spec.structure.separator_block_index, 0)
        self.assertGreaterEqual(spec.structure.section_for("summary").template_block_index, 0)

    def test_yck_ftr2_extracts_compact_margins_allcaps_no_separator(self) -> None:
        self.assertTrue(YCK_FTR2.is_file())
        spec = extract_format_from_docx(YCK_FTR2, name="YcKFTR2")
        self.assertAlmostEqual(spec.margin_top_in, 0.0, places=2)
        self.assertAlmostEqual(spec.margin_left_in, 0.4444, places=3)
        self.assertEqual(spec.name_size, 24)

        headings = {s.key: s.heading for s in spec.structure.sections}
        self.assertEqual(headings.get("summary"), "PROFESSIONAL SUMMARY")
        self.assertEqual(headings.get("experience"), "WORK EXPERIENCE")
        self.assertFalse(spec.structure.separator.enabled)
        self.assertEqual(spec.structure.experience.company_line_mode, "tabbed_line")
        self.assertEqual(spec.structure.experience.tab_pos_twips, 9000)
        self.assertEqual(spec.structure.contact_separator.strip(), "·")


class TemplatePreserveGenerateTests(unittest.TestCase):
    def test_upload_primary_generate_matches_yck_resume_not_default(self) -> None:
        from docx import Document
        from docx.oxml.ns import qn

        from resume_writer_app import build_resume, load_docx_dependencies, resume_format_from_store

        load_docx_dependencies()
        _ensure_user("tmplresume", "TmplResume1")
        saved = add_format_from_docx("tmplresume", YCK_RESUME, "YcK Resume Ref", set_as_primary=True)
        self.assertTrue(saved.template_file)
        template_path = resolve_format_template_path("tmplresume", saved)
        self.assertIsNotNone(template_path)
        self.assertTrue(template_path.is_file())
        self.assertEqual(get_primary_format("tmplresume").id, saved.id)

        fmt = resume_format_from_store("tmplresume")
        self.assertIsNotNone(fmt.template_path)
        self.assertTrue(Path(fmt.template_path).is_file())
        self.assertAlmostEqual(fmt.margin_top_in, 0.5, places=2)

        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "from_yck_resume.docx"
            build_resume(_sample_content(), fmt, out)
            doc = Document(str(out))
            section = doc.sections[0]
            self.assertAlmostEqual(section.top_margin.inches, 0.5, places=2)
            self.assertAlmostEqual(section.left_margin.inches, 0.5, places=2)
            # Must NOT be the protected default 0.75" layout.
            self.assertNotAlmostEqual(section.top_margin.inches, 0.75, places=2)

            texts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            self.assertIn("Professional Experience", texts)
            self.assertNotIn("WORK EXPERIENCE", texts)
            self.assertIn("Summary", texts)
            self.assertNotIn("SUMMARY", texts)

            # Separator prototype cloned from template (sz=4 black bottom border).
            found_sep = False
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    continue
                p_pr = paragraph._p.pPr
                if p_pr is None:
                    continue
                p_bdr = p_pr.find(qn("w:pBdr"))
                if p_bdr is None:
                    continue
                bottom = p_bdr.find(qn("w:bottom"))
                if bottom is None:
                    continue
                self.assertEqual(bottom.get(qn("w:sz")), "4")
                self.assertEqual(bottom.get(qn("w:color")), "000000")
                found_sep = True
                break
            self.assertTrue(found_sep, "expected cloned horizontal separator from template")

            # Experience uses tabbed lines (no two-column table), matching YcKResume.
            self.assertEqual(len(doc.tables), 0)
            dated = [p for p in doc.paragraphs if "\t" in p.text and "2020" in p.text]
            self.assertTrue(dated)
            tabs = dated[0]._p.pPr.find(qn("w:tabs")) if dated[0]._p.pPr is not None else None
            self.assertIsNotNone(tabs)

    def test_upload_primary_generate_matches_yck_ftr2_not_default(self) -> None:
        from docx import Document

        from resume_writer_app import build_resume, load_docx_dependencies, resume_format_from_store

        load_docx_dependencies()
        _ensure_user("tmplftr2", "TmplFtr21")
        saved = add_format_from_docx("tmplftr2", YCK_FTR2, "YcK FTR2 Ref", set_as_primary=True)
        self.assertTrue(saved.template_file)
        set_primary_format("tmplftr2", saved.id)

        fmt = resume_format_from_store("tmplftr2")
        self.assertIsNotNone(fmt.template_path)
        self.assertAlmostEqual(fmt.margin_top_in, 0.0, places=2)
        self.assertEqual(fmt.name_size, 24)

        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "from_ftr2.docx"
            build_resume(_sample_content(), fmt, out)
            doc = Document(str(out))
            section = doc.sections[0]
            self.assertAlmostEqual(section.top_margin.inches, 0.0, places=2)
            self.assertAlmostEqual(section.left_margin.inches, 0.4444, places=3)
            self.assertNotAlmostEqual(section.top_margin.inches, 0.75, places=2)

            texts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            self.assertIn("PROFESSIONAL SUMMARY", texts)
            self.assertIn("WORK EXPERIENCE", texts)
            self.assertNotIn("SUMMARY", texts)

            # FTR2 has no separator borders between sections.
            from docx.oxml.ns import qn

            sep_count = 0
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    continue
                p_pr = paragraph._p.pPr
                if p_pr is None:
                    continue
                p_bdr = p_pr.find(qn("w:pBdr"))
                if p_bdr is not None and p_bdr.find(qn("w:bottom")) is not None:
                    sep_count += 1
            self.assertEqual(sep_count, 0)

    def test_default_format_still_uses_historical_path_without_template(self) -> None:
        from resume_writer_app import ResumeFormat, build_resume, load_docx_dependencies

        load_docx_dependencies()
        from docx import Document

        fmt = ResumeFormat()  # protected default — no template_path
        self.assertIsNone(fmt.template_path)
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "default.docx"
            build_resume(_sample_content(), fmt, out)
            doc = Document(str(out))
            self.assertAlmostEqual(doc.sections[0].top_margin.inches, 0.75, places=2)
            texts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            self.assertIn("SUMMARY", texts)
            self.assertIn("WORK EXPERIENCE", texts)


if __name__ == "__main__":
    unittest.main()
